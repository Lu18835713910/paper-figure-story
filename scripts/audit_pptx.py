#!/usr/bin/env python3
"""Structural PowerPoint audit; requires rendering and scientific checking separately."""
import argparse
from pathlib import Path
from pptx import Presentation


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('pptx',type=Path)
    p.add_argument('--strict-images',action='store_true',help='Figure页必须至少含1张图片')
    p.add_argument('--no-logo-default',action='store_true',help='检查默认PFS无用户图标/图像')
    args=p.parse_args()
    prs=Presentation(str(args.pptx))
    errors,warns=[],[]
    if not prs.slides: errors.append('空PPTX')
    for i,s in enumerate(prs.slides,1):
        titles=[v.text.strip() for v in s.shapes if v.name=='PFS_TITLE' and v.has_text_frame]
        if not any(titles): errors.append(f'Slide {i} 缺少标题')
        if len(titles)>1: warns.append(f'Slide {i} 存在多个标题槽位')
        hasfig=any(v.name=='PFS_IMAGE' for v in s.shapes)
        images=[v for v in s.shapes if v.shape_type==13]
        if hasfig and not images:
            (errors if args.strict_images else warns).append(f'Slide {i} Figure图像缺失')
        for v in s.shapes:
            if v.name in ('PFS_IMAGE_PLACEHOLDER',) and v.has_text_frame and 'MISSING' in v.text:
                (errors if args.strict_images else warns).append(f'Slide {i} 仍含草稿占位字')
            if v.left < -15240 or v.top < -15240 or v.left+v.width > prs.slide_width+15240 or v.top+v.height > prs.slide_height+15240:
                warns.append(f'Slide {i} 形状越出画布：{v.name}')
        if args.no_logo_default and images and not hasfig:
            warns.append(f'Slide {i} 存在非Figure图片，应检查是否误加Logo')
    print('PPT结构检查：',args.pptx, '|',len(prs.slides),'slides')
    for w in warns: print('WARN:',w)
    for e in errors: print('ERROR:',e)
    if not errors and not warns: print('PASS（结构检查通过；仍需要目视检查排版与科学内容）')
    return 1 if errors else 0

if __name__=='__main__':raise SystemExit(main())
