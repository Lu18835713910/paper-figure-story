"""Visual Polish: bounded, source-preserving post-processing of an existing PPTX.

No PDF parsing, model inference, paper-specific rule, or network mutation occurs here.
Object diagnostics are candidates; an Agent must inspect actual PNGs.  Scientific
coverage gaps are handed upstream.  Only explicit, source-bound edits are applied.
"""
from __future__ import annotations
import argparse,copy,hashlib,io,json,os,posixpath,re,shutil,subprocess,sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR,MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from apply_script_fonts import apply_fonts,font_availability
from audit_script_fonts import check as check_fonts

VERSION='pfs-visual-polish-locked-2'
HERE=Path(__file__).resolve().parent
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
REL='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types'
MODES={'quick':1280,'balanced':1920,'full':2560}
class PolishError(ValueError):pass
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def fingerprint(x):return hashlib.sha256(json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def save(p,x):p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n','utf-8')
def normal(x):return re.sub(r'\s+','',str(x)).replace('•','')
def package(path):
 with ZipFile(path) as z:
  if z.testzip():raise PolishError('CORRUPT_PPTX')
  return {n:z.read(n) for n in z.namelist()}
def xml(data,part):return E.fromstring(data[part])
def put(data,part,root):data[part]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
def relpart(part):return posixpath.dirname(part)+'/_rels/'+posixpath.basename(part)+'.rels'

def resolve_target(part,target):return posixpath.normpath(target if target.startswith('/') else posixpath.dirname(part)+'/'+target).lstrip('/')
def active_parts(data):
 rels={r.get('Id'):r.get('Target') for r in xml(data,'ppt/_rels/presentation.xml.rels')}
 return [posixpath.normpath(t if t.startswith('/') else 'ppt/'+t).lstrip('/') for x in xml(data,'ppt/presentation.xml').findall('p:sldIdLst/p:sldId',NS) for t in [rels[x.get('{'+NS['r']+'}id')]]]
def shape_text(shape):return ''.join(shape.xpath('.//a:t/text()',namespaces=NS))
def texts(root):return [''.join(p.xpath('.//a:t/text()',namespaces=NS)) for p in root.xpath('.//a:p',namespaces=NS)]
def shape_id(shape):
 n=shape.find('.//p:cNvPr',NS);return int(n.get('id')) if n is not None else None
def shape_name(shape):
 n=shape.find('.//p:cNvPr',NS);return n.get('name','') if n is not None else ''

def declared_panels(text):
 # A literal panel claim is a requirement to check, never proof that the panel
 # exists. Without a source map the result is UNRESOLVED, not fabricated OCR.
 out=[]
 for m in re.finditer(r'(?:Figure|Fig\.?)\s*(\d+)\s*([A-Ha-h])(?:\s*[–—-]\s*([A-Ha-h]))?\b',text,re.I):
  first=ord(m[2].upper());last=ord((m[3] or m[2]).upper())
  if first>last:continue
  for letter in range(first,last+1):
   panel=m[1]+chr(letter)
   if panel not in out:out.append(panel)
 return out
def top_shape(root,sid):
 for s in root.find('p:cSld/p:spTree',NS):
  if shape_id(s)==int(sid):return s
 raise PolishError('SHAPE_NOT_TOP_LEVEL: '+str(sid))
def notes_part(data,part):
 r=relpart(part)
 if r not in data:return None
 for e in xml(data,r):
  if e.get('Type','').endswith('/notesSlide'):return resolve_target(part,e.get('Target'))
 return None
def whole_text(data,part):
 tx=texts(xml(data,part));n=notes_part(data,part)
 if n and n in data:tx+=texts(xml(data,n))
 return '\n'.join(tx)
def audit_input(path,context=None):
 p=Presentation(path);data=package(path);parts=active_parts(data);context=context or {};hints={int(x['slide']):x for x in context.get('slides',[])};records=[]
 for i,s in enumerate(p.slides,1):
  hint=hints.get(i,{});issues=[];objects=[];images=[];has_table=False
  for q in s.shapes:
   pos=[round(v/914400,5) for v in [q.left,q.top,q.width,q.height]]
   if min(pos[:2])<-.02 or q.left+q.width>p.slide_width+18288 or q.top+q.height>p.slide_height+18288:issues.append(dict(category='Layout',code='OFF_CANVAS',shape=q.shape_id,status='REVIEW_REQUIRED'))
   has_table |= q.has_table
   if q.has_text_frame and q.text.strip():
    sizes=[r.font.size.pt for pa in q.text_frame.paragraphs for r in pa.runs if r.font.size]
    objects.append(dict(id=q.shape_id,name=q.name,text=q.text,position=pos,font_sizes=sizes))
    if any(re.fullmatch(r'\s*[，。；：？！,.!?;:]\s*',ln) for ln in q.text.splitlines()):issues.append(dict(category='Typography',code='ISOLATED_PUNCTUATION',shape=q.shape_id,status='REVIEW_REQUIRED'))
   if q.shape_type==13:
    w,h=q.image.size;visible=[q.crop_left,q.crop_top,1-q.crop_right,1-q.crop_bottom];ratio=(w*(visible[2]-visible[0]))/(h*(visible[3]-visible[1]));error=abs(q.width/q.height/ratio-1)
    rec=dict(id=q.shape_id,sha256=hashlib.sha256(q.image.blob).hexdigest(),position=pos,pixels=[w,h],crop=visible,aspect_error=error);images.append(rec)
    if error>.025:issues.append(dict(category='Figure Integrity',code='IMAGE_ASPECT',shape=q.shape_id,status='REVIEW_REQUIRED'))
  body=[x for x in objects if x['position'][1]>.9 and x['position'][1]+x['position'][3]<p.slide_height/914400-.28]
  area=sum(x['position'][2]*x['position'][3] for x in body)+sum(x['position'][2]*x['position'][3] for x in images)
  text='\n'.join(x['text'] for x in objects)
  compare=hint.get('comparison') or (not images and text.count('•')>=2 and re.search(r'比较|对比|versus|\bvs\b|六个系统|三组',text,re.I))
  if compare and not has_table:issues.append(dict(category='Visual Encoding',code='COMPARISON_IN_LIST',status='REVIEW_REQUIRED',detail='Candidate for table/columns, not permission to invent missing cells.'))
  if area/(p.slide_width*p.slide_height/914400**2)<.20 and not images:issues.append(dict(category='Layout',code='LOW_CONTENT_OCCUPANCY',status='REVIEW_REQUIRED',detail='Whitespace is a review hint, not a quality verdict.'))
  for req in hint.get('required_content',[]):
   if not all(normal(t) in normal(text) for t in req.get('phrases',[])):issues.append(dict(category='Content Coverage',code='UPSTREAM_CONTENT_REQUIRED',request_id=req['id'],status='REVIEW_REQUIRED',detail=req.get('description','Missing comparison content; return to planning.')))
  mapped=[]
  for im in images:
   asset=next((a for a in context.get('assets',[]) if a.get('sha256')==im['sha256']),None)
   if asset is None:continue
   bbox=asset.get('bbox_px',[0,0,*im['pixels']]);x,y,xx,yy=bbox;v=im['crop'];visible=[x+(xx-x)*v[0],y+(yy-y)*v[1],x+(xx-x)*v[2],y+(yy-y)*v[3]]
   for panel in asset.get('panels',[]) if asset.get('panel_mapping_status')=='CHECKED' else []:
    if panel!='all':mapped.append(panel)
   for region in asset.get('critical_regions',[]):
    a,b,c,d=region['bbox_px']
    if not (visible[0]<=a and visible[1]<=b and visible[2]>=c and visible[3]>=d):issues.append(dict(category='Figure Integrity',code='CRITICAL_REGION_CLIPPED',shape=im['id'],region=region['id'],status='REVIEW_REQUIRED'))
  expected=list(dict.fromkeys(hint.get('expected_panels',[])+declared_panels(text)))
  missing=[x for x in expected if x not in mapped]
  if missing:issues.append(dict(category='Figure Integrity',code='PANEL_MAPPING_UNRESOLVED',panels=missing,status='REVIEW_REQUIRED',detail='Missing or not registered; no OCR/semantic verification is claimed.'))
  if hint.get('scientific_image') and not images:issues.append(dict(category='Figure Integrity',code='MISSING_FIGURE',status='REVIEW_REQUIRED'))
  title=next((x['text'] for x in objects if x['name']=='PFS_TITLE'),None) or max(objects,key=lambda x:max(x['font_sizes'] or [0]),default={'text':''})['text']
  records.append(dict(slide=i,part=parts[i-1],title=title,objects=objects,images=images,native_table=has_table,issues=issues,source_status=context.get('scientific_review','NOT_AVAILABLE'),science_review='NOT_PERFORMED_BY_VISUAL_POLISH'))
 return dict(version=VERSION,input=str(Path(path).resolve()),input_sha256=digest(path),slide_size=[p.slide_width/914400,p.slide_height/914400],slides=records,visual_review='PENDING_ACTUAL_IMAGE_REVIEW')

def render(path,folder,profile,slides=None,provided=None):
 folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
 if provided:
  receipt=read(provided)
  if receipt.get('input_sha256')!=digest(path) or not receipt.get('source_unchanged'):raise PolishError('PREVIEW_RECEIPT_DRIFT')
  if not receipt.get('rendered'):raise PolishError('EMPTY_RENDER_RECEIPT')
  for r in receipt['rendered']:
   if digest(r['png'])!=r['png_sha256']:raise PolishError('PREVIEW_IMAGE_DRIFT')
  save(folder/'receipt.json',receipt);return receipt
 if os.name!='nt':raise PolishError('POWERPOINT_RENDERER_UNAVAILABLE: provide a fingerprinted renderer receipt')
 exe=shutil.which('powershell.exe') or shutil.which('pwsh.exe')
 if not exe:raise PolishError('POWERSHELL_UNAVAILABLE')
 receipt=folder/'receipt.json';spec=dict(input=str(Path(path).resolve()),input_sha256=digest(path),outdir=str(folder.resolve()),receipt=str(receipt.resolve()),width=MODES[profile])
 if slides:spec['slides']=slides
 f=folder/'render_manifest.json';save(f,spec)
 r=subprocess.run([exe,'-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(HERE/'visual_polish_office.ps1'),'-Manifest',str(f)],capture_output=True,text=True,encoding='utf-8',errors='replace')
 save(folder/'render_execution.json',dict(command=[exe,str(HERE/'visual_polish_office.ps1'),str(f)],exit_code=r.returncode,stdout=r.stdout,stderr=r.stderr))
 if r.returncode or not receipt.is_file():raise PolishError('POWERPOINT_RENDER_FAILED: '+str(folder/'render_execution.json'))
 result=read(receipt)
 if not result.get('rendered'):raise PolishError('EMPTY_RENDER_RECEIPT')
 return result

def verify_context(context):
 protected=context.get('protected_sources',[])
 for r in protected:
  if digest(r['path'])!=r['sha256']:raise PolishError('PROTECTED_SOURCE_DRIFT: '+r['path'])
 for r in context.get('sources',[]):
  if digest(r['path'])!=r['sha256']:raise PolishError('DISPLAY_SOURCE_DRIFT: '+r['id'])
 if context.get('template'):
  t=context['template']
  if Path(t['path']).suffix.lower()!='.pptx' or digest(t['path'])!=t['sha256']:raise PolishError('TEMPLATE_REFERENCE_DRIFT')
  package(t['path'])
 return {r['path']:r['sha256'] for r in protected}
def source_text(context,rid):
 r=next((x for x in context.get('sources',[]) if x['id']==rid),None)
 if r is None:raise PolishError('UNKNOWN_DISPLAY_SOURCE: '+rid)
 text=Path(r['path']).read_text(encoding='utf-8-sig')
 return text
def bound_text(record,original,parts,context):
 text=record['text'] if isinstance(record,dict) else str(record)
 refs=record.get('refs',[]) if isinstance(record,dict) else []
 if not refs:raise PolishError('UNBOUND_TEXT')
 ok=False
 for ref in refs:
  if ref.get('kind')=='deck':
   n=int(ref['slide']);pool=whole_text(original,parts[n-1]);ok |= normal(text) in normal(pool)
  elif ref.get('kind')=='upstream':
   # Upstream must supply the text; this module never writes new scientific cells.
   r=next((x for x in context.get('upstream_responses',[]) if x['id']==ref['response']),None)
   if not r or not r.get('request_id') or r.get('source_review')!='AGENT_CHECKED_PARTIAL':raise PolishError('UPSTREAM_NOT_ADJUDICATED')
   payload=read(r['path'])
   if digest(r['path'])!=r['sha256'] or not payload.get('source_bindings'):raise PolishError('UPSTREAM_RESPONSE_DRIFT_OR_UNSOURCED')
   for binding in payload['source_bindings']:
    if digest(binding['path'])!=binding['sha256']:raise PolishError('UPSTREAM_BINDING_DRIFT')
    if binding.get('excerpt') and binding['excerpt'] not in Path(binding['path']).read_text('utf-8-sig'):raise PolishError('UPSTREAM_EXCERPT')
   ok |= normal(text) in normal(json.dumps(payload,ensure_ascii=False))
  else:raise PolishError('UNKNOWN_TEXT_BINDING')
 if not ok:raise PolishError('TEXT_NOT_IN_BOUND_SOURCE: '+text[:80])
 return text
def shape_xml(kind,text,pos,size=16,bold=False,color='17345E'):
 prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);q=s.shapes.add_textbox(*[Inches(v) for v in pos]);q.name=kind
 tf=q.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.03);tf.margin_top=tf.margin_bottom=Inches(.015)
 for i,ln in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'));p.space_after=Pt(0)
 return copy.deepcopy(q._element)
def native_table(headers,rows,pos,widths,size):
 if size<13:raise PolishError('TABLE_FONT_TOO_SMALL')
 prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);q=s.shapes.add_table(len(rows)+1,len(headers),*[Inches(v) for v in pos]);q.name='VP_COMPARISON_TABLE';t=q.table
 if widths:
  if abs(sum(widths)-pos[2])>.02:raise PolishError('TABLE_WIDTHS')
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate([headers]+rows):
  for j,txt in enumerate(row):
   c=t.cell(i,j);c.text=txt;c.margin_left=c.margin_right=Inches(.06);c.margin_top=c.margin_bottom=Inches(.04);c.vertical_anchor=MSO_ANCHOR.MIDDLE
   c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string('174994' if i==0 else ('F3F6FA' if i%2 else 'FFFFFF'))
   for p in c.text_frame.paragraphs:
    p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '17345E');p.space_after=Pt(0)
 return copy.deepcopy(q._element)
def append_shape(root,shape):
 shape=E.fromstring(E.tostring(shape))
 maxid=max([shape_id(x) or 0 for x in root.find('p:cSld/p:spTree',NS)],default=1)
 for n in shape.xpath('.//p:cNvPr',namespaces=NS):n.set('id',str(maxid+1))
 root.find('p:cSld/p:spTree',NS).append(shape)
def delete_shapes(root,ids):
 for sid in ids:
  if root.xpath('.//p:spTgt[@spid="'+str(sid)+'"]',namespaces=NS):raise PolishError('ANIMATED_SHAPE_REQUIRES_REVIEW')
  s=top_shape(root,sid);root.find('p:cSld/p:spTree',NS).remove(s)
def add_note(data,part,text):
 n=notes_part(data,part)
 if not n:raise PolishError('NOTES_PART_REQUIRED_FOR_REWRITE')
 root=xml(data,n);body=next((s for s in root.findall('p:cSld/p:spTree/p:sp',NS) if s.find('.//p:ph[@type="body"]',NS) is not None),None)
 if body is None:raise PolishError('NOTES_BODY_NOT_FOUND')
 tx=body.find('p:txBody',NS);p=E.SubElement(tx,'{'+NS['a']+'}p');r=E.SubElement(p,'{'+NS['a']+'}r');E.SubElement(r,'{'+NS['a']+'}t').text=text;put(data,n,root)
def insert_slide(data,source,partname,after,baseline=None):
 # Clone only the chosen style/relationships; original slide parts remain intact.
 baseline=baseline or data
 data[partname]=baseline[source];rp=relpart(source);nr=E.fromstring(baseline[rp])
 for r in list(nr):
  if r.get('Type','').endswith('/notesSlide'):
   old=resolve_target(source,r.get('Target'));new='ppt/notesSlides/'+Path(partname).stem+'.xml';data[new]=baseline[old]
   nrels=relpart(old)
   if nrels in data:
    rr=xml(baseline,nrels)
    for e in rr:
     if e.get('Type','').endswith('/slide'):e.set('Target','../slides/'+posixpath.basename(partname))
    put(data,relpart(new),rr)
   r.set('Target','../notesSlides/'+posixpath.basename(new))
   types=xml(data,'[Content_Types].xml');E.SubElement(types,'{'+CT+'}Override',PartName='/'+new,ContentType='application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml');put(data,'[Content_Types].xml',types)
 put(data,relpart(partname),nr)
 pr=xml(data,'ppt/presentation.xml');rels=xml(data,'ppt/_rels/presentation.xml.rels');rid='rId'+Path(partname).stem;E.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=NS['r']+'/slide',Target='slides/'+posixpath.basename(partname));put(data,'ppt/_rels/presentation.xml.rels',rels)
 ids=pr.find('p:sldIdLst',NS);new=E.Element('{'+NS['p']+'}sldId',id=str(max(int(x.get('id')) for x in ids)+1));new.set('{'+NS['r']+'}id',rid);ids.insert(after,new);put(data,'ppt/presentation.xml',pr)
 types=xml(data,'[Content_Types].xml');E.SubElement(types,'{'+CT+'}Override',PartName='/'+partname,ContentType='application/vnd.openxmlformats-officedocument.presentationml.slide+xml');put(data,'[Content_Types].xml',types)
 return partname
def mutate(data,original,parts,action,context,index):
 raise PolishError('LEGACY_CONTENT_MUTATION_DISABLED: Visual Polish is layout/style only')
 kind=action['kind'];part=parts[int(action['slide'])-1]
 if kind=='insert_table_slide':part=insert_slide(data,part,'ppt/slides/visualPolish'+str(index)+'.xml',int(action['after']),original)
 root=xml(data,part);removed=[]
 if kind in ['table','insert_table_slide']:
  rows=[[bound_text(c,original,parts,context) for c in row] for row in action['rows']]
  generic={'系统','模型','输入','输出/用途','任务','已完成的检验','已有证据','证据边界','仍需注意','研究问题','比较项目','当前结论','下一步','原文证据','已有结果','这说明什么','方法/读数','观察结果','科学意义','条件','原有取样','拟议设计','已知什么','为何继续','研究对象'}
  headers=[h if isinstance(h,str) and h in generic else bound_text(h,original,parts,context) for h in action['headers']]
  for sid in action.get('remove_shapes',[]):removed.append(shape_text(top_shape(root,sid)))
  if kind=='insert_table_slide':
   for s in list(root.find('p:cSld/p:spTree',NS)):
    if shape_name(s) in action.get('clear_names',[]):root.find('p:cSld/p:spTree',NS).remove(s)
  else:delete_shapes(root,action.get('remove_shapes',[]))
  append_shape(root,native_table(headers,rows,action['position'],action.get('column_widths'),action.get('font_size',15)))
 elif kind=='layout':
  for change in action['shapes']:
   s=top_shape(root,change['shape_id'])
   if s.tag=='{'+NS['p']+'}grpSp':raise PolishError('GROUP_REQUIRES_TARGETED_REVIEW')
   xfrm=s.find('.//a:xfrm',NS)
   if xfrm is None:raise PolishError('GROUP_OR_UNKNOWN_TRANSFORM')
   x,y,w,h=change['position'];xfrm.find('a:off',NS).attrib.update({'x':str(round(x*914400)),'y':str(round(y*914400))});xfrm.find('a:ext',NS).attrib.update({'cx':str(round(w*914400)),'cy':str(round(h*914400))})
   if change.get('font_size'):
    if change['font_size']<12:raise PolishError('BODY_FONT_TOO_SMALL')
    for p in s.xpath('.//a:rPr|.//a:defRPr|.//a:endParaRPr',namespaces=NS):p.set('sz',str(round(change['font_size']*100)))
 elif kind=='replace_text':
  s=top_shape(root,action['shape_id']);old=shape_text(s)
  if normal(action['expected_text'])!=normal(old):raise PolishError('TEXT_INPUT_DRIFT')
  new=bound_text(action['replacement'],original,parts,context)
  # Keep original detailed wording in existing speaker notes, never silently discard it.
  removed.append(old);delete_shapes(root,[action['shape_id']]);append_shape(root,shape_xml('VP_READING_TEXT',new,action['position'],action.get('font_size',17),action.get('bold',False)))
 elif kind=='add_text':
  new=bound_text(action['content'],original,parts,context);append_shape(root,shape_xml('VP_BOUND_CALLOUT',new,action['position'],action.get('font_size',15),action.get('bold',False)))
 elif kind=='narrative':
  if action.get('identity')!='presentation':
   edge=next((e for e in context.get('relationships',[]) if e['id']==action.get('relation_id')),None)
   if edge is None:raise PolishError('UNREGISTERED_SCIENTIFIC_ARROW')
   if action.get('identity')!=edge.get('identity'):raise PolishError('ARROW_IDENTITY_UPGRADE')
  for sid in action.get('remove_shapes',[]):removed.append(shape_text(top_shape(root,sid)))
  delete_shapes(root,action.get('remove_shapes',[]))
  for item in action['steps']:
   append_shape(root,shape_xml('VP_NARRATIVE',bound_text(item,original,parts,context),item['position'],item.get('font_size',17),item.get('bold',False)))
  for a,b in action.get('arrows',[]):
   prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);q=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));q.name='VP_PRESENTATION_ARROW';q.line.width=Pt(1.8);q.line.color.rgb=RGBColor.from_string('174994');ln=q._element.spPr.find('{'+NS['a']+'}ln');E.SubElement(ln,'{'+NS['a']+'}tailEnd',type='triangle',w='lg',len='lg');append_shape(root,copy.deepcopy(q._element))
 elif kind=='replace_image':
  asset=next((a for a in context.get('assets',[]) if a['id']==action['asset_id']),None)
  if not asset or digest(asset['path'])!=asset['sha256'] or asset.get('panel_mapping_status')!='CHECKED' or asset.get('visual_status')!='SOURCE_CHECKED':raise PolishError('IMAGE_NOT_SOURCE_REVIEWED')
  if 'all' in action.get('required_panels',[]) or not set(action.get('required_panels',[]))<=set(asset.get('panels',[])):raise PolishError('IMAGE_PANEL_MAPPING')
  s=top_shape(root,action['shape_id']);blip=s.find('.//a:blip',NS)
  if blip is None:raise PolishError('NOT_A_PICTURE')
  target='ppt/media/visualPolish'+asset['sha256'][:24]+Path(asset['path']).suffix.lower();data[target]=Path(asset['path']).read_bytes();rp=relpart(part);rels=xml(data,rp);rid='rIdVPImage'+str(index);E.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=NS['r']+'/image',Target='../media/'+posixpath.basename(target));put(data,rp,rels);blip.set('{'+NS['r']+'}embed',rid)
  fill=s.find('p:blipFill',NS);crop=fill.find('a:srcRect',NS)
  if crop is not None:fill.remove(crop)
  x,y,w,h=action['position'];iw,ih=Image.open(asset['path']).size;ww=min(w,h*iw/ih);hh=ww*ih/iw;x+=(w-ww)/2;y+=(h-hh)/2;tr=s.find('.//a:xfrm',NS);tr.find('a:off',NS).attrib.update({'x':str(round(x*914400)),'y':str(round(y*914400))});tr.find('a:ext',NS).attrib.update({'cx':str(round(ww*914400)),'cy':str(round(hh*914400))})
  types=xml(data,'[Content_Types].xml');ext=Path(target).suffix[1:]
  if not any(e.get('Extension')==ext for e in types):E.SubElement(types,'{'+CT+'}Default',Extension=ext,ContentType='image/'+('jpeg' if ext in ['jpg','jpeg'] else ext));put(data,'[Content_Types].xml',types)
 elif kind=='backup':
  # Never hide an important limitation unless the exact text remains in the main flow.
  requirements=list(action.get('keep_visible',[]))
  requirements+=next((h.get('important_limits',[]) for h in context.get('slides',[]) if h['slide']==action['slide']),[])
  for required in requirements:
   phrase=bound_text(required,original,parts,context)
   other=[p for p in active_parts(data) if p!=part and xml(data,p).get('show')!='0']
   if not any(normal(phrase) in normal('\n'.join(texts(xml(data,p)))) for p in other):raise PolishError('SCIENTIFIC_LIMIT_WOULD_BE_HIDDEN')
  root.set('show','0')
  if action.get('move_to_end',True):
   pr=xml(data,'ppt/presentation.xml');rels={r.get('Id'):posixpath.normpath('ppt/'+r.get('Target')) for r in xml(data,'ppt/_rels/presentation.xml.rels')};ids=pr.find('p:sldIdLst',NS);node=next(v for v in ids if rels[v.get('{'+NS['r']+'}id')]==part);ids.remove(node);ids.append(node);put(data,'ppt/presentation.xml',pr)
 else:raise PolishError('UNSUPPORTED_POLISH_ACTION: '+kind)
 if action.get('title'):
  s=next((x for x in root.find('p:cSld/p:spTree',NS) if shape_name(x)=='PFS_TITLE'),None)
  if s is not None:
   title=bound_text(action['title'],original,parts,context);old=shape_text(s);root.find('p:cSld/p:spTree',NS).remove(s);append_shape(root,shape_xml('PFS_TITLE',title,action.get('title_position',[.72,.67,11.93,.75]),action.get('title_font_size',26),True))
   removed.append(old)
 put(data,part,root)
 if removed:add_note(data,part,'Visual Polish保留的原页文字（用于复核，不新增科研事实）：\n'+'\n'.join(t for t in removed if t))
 return part

def before_after(before,after,out):
 a=Image.open(before).convert('RGB');b=Image.open(after).convert('RGB');a.thumbnail((1000,760));b.thumbnail((1000,760));im=Image.new('RGB',(2040,820),'#EDF2F8');draw=ImageDraw.Draw(im)
 try:font=ImageFont.truetype('C:/Windows/Fonts/times.ttf',22)
 except OSError:font=ImageFont.load_default()
 draw.text((20,12),'Before',fill='#174994',font=font);draw.text((1040,12),'After - user review pending',fill='#174994',font=font);im.paste(a,(20,50));im.paste(b,(1040,50));im.save(out)

def _legacy_polish(input_path,outdir,profile='balanced',context=None,plan=None,before_receipt=None,renderer='powerpoint',cycle=1):
 raise PolishError('LEGACY_CONTENT_MUTATION_DISABLED: use the content-locked v2 entry')
 src=Path(input_path).resolve();out=Path(outdir).resolve()
 if profile not in MODES or cycle not in [1,2]:raise PolishError('PROFILE_OR_CYCLE')
 if out.exists() and any(out.iterdir()):raise PolishError('OUTPUT_EXISTS: preserve earlier run; use a new output directory')
 if src.suffix.lower()!='.pptx' or not src.is_file():raise PolishError('PPTX_INPUT_REQUIRED')
 if src==out/'polished.pptx':raise PolishError('INPUT_OVERWRITE_FORBIDDEN')
 context=read(context) if isinstance(context,(str,Path)) else context or {};plan=read(plan) if isinstance(plan,(str,Path)) else plan or {};protected=verify_context(context)
 if plan and (plan.get('version')!=VERSION or plan.get('input_sha256')!=digest(src)):raise PolishError('PLAN_VERSION_OR_INPUT_DRIFT')
 out.mkdir(parents=True,exist_ok=True);shutil.copy2(src,out/'original.pptx');audit=audit_input(src,context);save(out/'diagnosis.json',audit)
 before=None
 if renderer!='none':before=render(src,out/'previews/before',profile,provided=before_receipt)
 else:save(out/'render_not_performed.json',dict(status='NOT_RENDERED',reason='Explicit renderer=none; no visual pass is possible.'))
 bybefore={r['slide']:r for r in (before or {}).get('rendered',[])}
 for r in plan.get('visual_reviews',[]):
  if not before or int(r['slide']) not in bybefore or bybefore[int(r['slide'])]['png_sha256']!=r.get('preview_sha256'):raise PolishError('VISUAL_REVIEW_PREVIEW_DRIFT')
 for page in audit['slides']:
  for r in plan.get('visual_reviews',[]):
   if int(r['slide'])==page['slide']:page['issues']+=r.get('findings',[])
 save(out/'diagnosis.json',audit)
 requests=[dict(slide=p['slide'],**x) for p in audit['slides'] for x in p['issues'] if x['category']=='Content Coverage']
 save(out/'upstream_requests.json',dict(version=VERSION,requests=requests,rule='Content Coverage is not invented by Visual Polish; explicit source-bound planning response is required.'))
 data=package(src);original=copy.deepcopy(data);parts=active_parts(data);actions=[];failed=[]
 for i,a in enumerate(plan.get('actions',[]),1):
  trial=copy.deepcopy(data)
  try:
   target=mutate(trial,original,parts,a,context,i);data=trial;actions.append(dict(id=a['id'],kind=a['kind'],original_slide=a['slide'],part=target,status='APPLIED',upstream_response=a.get('upstream_response')))
  except (PolishError,KeyError,ValueError,IndexError,TypeError) as e:failed.append(dict(id=a.get('id'),slide=a.get('slide'),kind=a.get('kind'),status='REVIEW_REQUIRED_OR_BLOCKED',reason=str(e),original_page_retained=True))
 # Targeted slide edits are merged into the original ZIP.  Charts, media, OLE,
 # masters, unrelated XML and source identifiers are not rebuilt by python-pptx.
 final=out/'polished.pptx'
 with ZipFile(final,'w',ZIP_DEFLATED) as z:
  for n,v in data.items():z.writestr(n,v)
 font_receipt=apply_fonts(final);save(out/'font_policy.json',font_receipt);fa=check_fonts(final);save(out/'font_audit.json',fa)
 newdata=package(final);newparts=active_parts(newdata)
 # Page counters are navigation metadata, not scientific wording. Reordering
 # never silently leaves duplicated source page numbers in the final main flow.
 for i,part in enumerate(newparts,1):
  root=xml(newdata,part)
  for q in root.find('p:cSld/p:spTree',NS):
   if shape_name(q)=='PFS_NUMBER':
    runs=q.xpath('.//a:t',namespaces=NS)
    for j,t in enumerate(runs):t.text=f'{i:02d}' if j==0 else ''
  put(newdata,part,root)
 with ZipFile(final,'w',ZIP_DEFLATED) as z:
  for n,v in newdata.items():z.writestr(n,v)
 changed=[i for i,p in enumerate(newparts,1) if original.get(p)!=newdata[p]]
 preserved_media=all(n in newdata and newdata[n]==v for n,v in original.items() if n.startswith('ppt/media/'))
 preserved_sources=all(digest(p)==h for p,h in protected.items())
 # Every removed literal is retained in target notes; positional/font-only edits
 # leave words unchanged.  This is not a semantic proof for upstream translations.
 before_lines=[normal(t) for part in parts for q in xml(original,part).find('p:cSld/p:spTree',NS) if shape_name(q)!='PFS_NUMBER' for t in texts(q) if normal(t)]
 after_pool=normal('\n'.join(whole_text(newdata,p) for p in newparts));lost=[x for x in before_lines if x not in after_pool]
 invariants=dict(original_input_unchanged=digest(src)==audit['input_sha256'],protected_sources_unchanged=preserved_sources,original_media_parts_unchanged=preserved_media,unintended_text_loss=lost,formal_science_status_unchanged=True,scientific_edges_added=0)
 save(out/'invariants.json',invariants)
 after=None
 if renderer!='none':
  after=render(final,out/'previews/after',profile,slides=changed or [1])
  rendered={r['slide']:r for r in after['rendered']};(out/'before_after').mkdir()
  for i,p in enumerate(newparts,1):
   oldindex=parts.index(p)+1 if p in parts else None
   if i not in rendered and oldindex in bybefore:
    old=bybefore[oldindex];dest=out/'previews/after'/f'slide-{i:03d}.png';shutil.copy2(old['png'],dest);rendered[i]=dict(old,slide=i,png=str(dest),reused_unchanged_part=True)
   if i in changed and i in rendered and oldindex in bybefore:before_after(bybefore[oldindex]['png'],rendered[i]['png'],out/'before_after'/f'slide-{i:03d}.png')
  after['rendered']=list(sorted(rendered.values(),key=lambda r:r['slide']));save(out/'previews/after/receipt.json',after)
 final_context=copy.deepcopy(context)
 final_context['slides']=[dict(h,slide=newparts.index(parts[int(h['slide'])-1])+1) for h in context.get('slides',[])]
 post=audit_input(final,final_context);save(out/'post_diagnosis.json',post)
 auto_issues=[dict(slide=r['slide'],**v) for r in (after or {}).get('rendered',[]) for v in r['issues']]
 auto_issues += [dict(slide=p['slide'],**v) for p in post['slides'] for v in p['issues'] if v['code'] in {'OFF_CANVAS','IMAGE_ASPECT','CRITICAL_REGION_CLIPPED','MISSING_FIGURE','PANEL_MAPPING_UNRESOLVED','ISOLATED_PUNCTUATION'}]
 def resolved(page,issue):
  ids=plan.get('resolves',{}).get(str(page)+':'+issue['code'],[])
  performed=[a for a in actions if a['id'] in ids]
  if issue['category']=='Content Coverage':
   return any(any(r['id']==a.get('upstream_response') and r.get('request_id')==issue.get('request_id') for r in context.get('upstream_responses',[])) for a in performed)
  if performed:return True
  # A real image review may retain intentional whitespace; never waive missing
  # panels, clipped regions, unknown science, overflow or missing content this way.
  if issue['code'] not in {'LOW_CONTENT_OCCUPANCY','COMPARISON_IN_LIST'}:return False
  return any(int(r['slide'])==page and any(d.get('code')==issue['code'] and d.get('disposition')=='RETAIN_WITH_REASON' and d.get('reason') for d in r.get('decisions',[])) for r in plan.get('visual_reviews',[]))
 unresolved=[dict(slide=p['slide'],**i) for p in audit['slides'] for i in p['issues'] if not resolved(p['slide'],i)]
 checks_ok=all([invariants['original_input_unchanged'],preserved_sources,preserved_media,not lost,fa['valid'],not auto_issues,not failed])
 report=dict(version=VERSION,input=str(src),input_sha256=audit['input_sha256'],output=str(final),output_sha256=digest(final),profile=profile,cycle=cycle,slides_before=len(parts),slides_after=len(newparts),actions=actions,failed_actions=failed,unresolved=unresolved,upstream_requests=requests,invariants=invariants,automatic_checks='PASS' if checks_ok else 'PARTIAL',actual_powerpoint_export=after is not None,automatic_issues=auto_issues,changed_pages=changed,agent_visual_review='PENDING_FINAL_IMAGES',USER_ACCEPTANCE='PENDING',scientific_review=context.get('scientific_review','NOT_AVAILABLE'),science_audit='NOT_PERFORMED_BY_VISUAL_POLISH',status='POLISHED_PENDING_VISUAL_REVIEW' if checks_ok and after else 'PARTIAL',source_context_fingerprint=fingerprint(context),render_width=MODES[profile],scientific_model_reads='Read-only fingerprints/context only; no extraction or inference',cost_policy='One before render, one affected-page render; no internal optimization loop; maximum cycle=2.')
 save(out/'visual_polish_report.json',report);write_report(out,report)
 return report
def write_report(out,r):
 lines=['# Visual Polish report','',f"状态：{r['status']}；自动检查：{r['automatic_checks']}；最终图像审阅：{r['agent_visual_review']}；USER_ACCEPTANCE=PENDING。",'',f"输入{r['slides_before']}页，输出{r['slides_after']}页；原稿保存在original.pptx。已应用{len(r['actions'])}项定向操作，未执行{len(r['failed_actions'])}项。",'','此模块不重新理解论文，不更改科研事实模型；对象/文件校验不等于科学终审。','', '## 应用记录']
 lines += [f"- {a['id']}：原第{a['original_slide']}页，{a['kind']}。" for a in r['actions']]
 lines += ['','## 不能安全修复/仍待核对']+[f"- 第{x.get('slide','?')}页：{x.get('code',x.get('reason'))}" for x in r['unresolved']+r['failed_actions']]
 boundary='仅调整原页已有信息的展示；页数、顺序、正文、备注、Figure归属和逻辑关系锁定；内容不足只报告UPSTREAM_CONTENT_ISSUE。' if r.get('version')=='pfs-visual-polish-locked-2' else '历史边界：上游内容绑定不构成科学终审。'
 lines += ['', '原图/来源保持校验：'+str(r['invariants'])+'。'+boundary,'前后预览在before_after/，完整最终预览在previews/after/；实际Office边界检查独立保存。']
 (Path(out)/'visual_polish_report.md').write_text('\n'.join(lines)+'\n','utf-8')
def _legacy_finish_review(report_path,decisions):
 p=Path(report_path);r=read(p);v=read(decisions) if isinstance(decisions,(str,Path)) else decisions;receipt=read(p.parent/'previews/after/receipt.json');pages={x['slide']:x for x in receipt['rendered']}
 if v.get('output_sha256')!=digest(r['output']):raise PolishError('FINAL_REVIEW_OUTPUT_DRIFT')
 needed=set(r['changed_pages']) or {1};reviewed=set()
 for x in v['pages']:
  i=int(x['slide'])
  if i not in pages or digest(pages[i]['png'])!=x['png_sha256'] or x['png_sha256']!=pages[i]['png_sha256']:raise PolishError('FINAL_REVIEW_IMAGE_DRIFT')
  if x['status']=='PASS':reviewed.add(i)
 # Only the encoding suggestion can be closed by observing a native table in
 # the reviewed final image. Content gaps and scientific/source failures cannot.
 post=read(p.parent/'post_diagnosis.json');tables={x['slide'] for x in post['slides'] if x['native_table']}
 resolved=[x for x in r['unresolved'] if x['code']=='COMPARISON_IN_LIST' and x['slide'] in tables & reviewed]
 r['resolved_in_final_visual_review']=resolved
 r['unresolved']=[x for x in r['unresolved'] if x not in resolved]
 r['agent_visual_review']='REVIEWED' if needed<=reviewed else 'PARTIAL';r['independent_user_or_reader_test']='NOT_PERFORMED';r['USER_ACCEPTANCE']='PENDING';r['visual_decisions']=v
 r['status']='POLISHED_AGENT_REVIEWED' if needed<=reviewed and r['automatic_checks']=='PASS' and not r['unresolved'] else 'PARTIAL'
 save(p,r);write_report(p.parent,r);return r
def polish(input_path,outdir,profile='balanced',context=None,plan=None,before_receipt=None,renderer='powerpoint',cycle=1):
 # The default and automatic PFS entry share the same content-locked boundary.
 # Legacy insert/backup/replace-image/upstream edit plans are intentionally refused.
 from visual_polish_locked import polish as locked_polish
 return locked_polish(input_path,outdir,profile,context,plan,before_receipt,renderer,cycle)

def finish_review(report_path,decisions):
 from visual_polish_locked import finish_review as locked_review
 return locked_review(report_path,decisions)

def main(argv=None):
 argv=list(sys.argv[1:] if argv is None else argv)
 if argv and argv[0]=='review':
  p=argparse.ArgumentParser(description='Record actual final PNG review without changing the PPTX');p.add_argument('--report',required=True);p.add_argument('--decisions',required=True);a=p.parse_args(argv[1:]);r=finish_review(a.report,a.decisions)
 else:
  p=argparse.ArgumentParser(description=__doc__);p.add_argument('--pptx',required=True);p.add_argument('--outdir',required=True);p.add_argument('--profile',choices=list(MODES),default='balanced');p.add_argument('--context');p.add_argument('--plan');p.add_argument('--before-receipt');p.add_argument('--renderer',choices=['powerpoint','none'],default='powerpoint');p.add_argument('--cycle',type=int,choices=[1,2],default=1);a=p.parse_args(argv);r=polish(a.pptx,a.outdir,a.profile,a.context,a.plan,a.before_receipt,a.renderer,a.cycle)
 print(json.dumps(dict(status=r['status'],output=r['output'],automatic_checks=r['automatic_checks'],USER_ACCEPTANCE='PENDING'),ensure_ascii=False));return 0 if r['automatic_checks']=='PASS' else 2
if __name__=='__main__':
 try:raise SystemExit(main())
 except (PolishError,OSError,KeyError,ValueError,TypeError) as e:print(str(e),file=sys.stderr);raise SystemExit(2)
