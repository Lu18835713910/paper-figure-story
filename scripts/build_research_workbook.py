"""Structured DOCX workbook from the same research-logic model and PPT exports.

Full scientific facts remain in M1/Inventory. This optional document contains
logic diagrams, source-linked explanations and structured reference appendices.
"""
from __future__ import annotations
import argparse,copy,csv,hashlib,json,shutil
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import research_logic as rl

def value(x):
    if x is None:return '待核实'
    return x if isinstance(x,str) else json.dumps(x,ensure_ascii=False)
def fonts(run,size=10.5,bold=False,color=None):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key,font in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','Microsoft YaHei')]:rf.set(qn('w:'+key),font)
    if color:run.font.color.rgb=RGBColor.from_string(color)
def table(doc,headers,rows,widths,size=10):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):
        c.text=h;pr=c._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),'174994');pr.append(s)
        for r in c.paragraphs[0].runs:fonts(r,size,True,'FFFFFF')
    for c,w in zip(t.rows[0].cells,widths):c.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,w in zip(cells,widths):c.width=Inches(w)
        for c,txt in zip(cells,row):
            c.text=value(txt)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
                for r in p.runs:fonts(r,size)
        props=cells[0]._tc.getparent().get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
    return t
def hyperlink(p,label,target):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'174994');rp.append(co);r.append(rp);tt=OxmlElement('w:t');tt.text=label;r.append(tt);h.append(r);p._p.append(h)
def page_setup(section,width=11.69,height=8.27):
    section.page_width=Inches(width);section.page_height=Inches(height)
    section.top_margin=section.bottom_margin=Inches(.45);section.left_margin=section.right_margin=Inches(.55)
def build(model_path,manifest_path,output,annotations_path=None):
    model=rl.load(model_path);check=rl.validate_model(model)
    if not check['valid']:raise ValueError(str(check))
    net,inv=rl.inputs(model);methods=rl.methods_view(model);em={x['experiment_id']:x for x in methods};nm={x['id']:x for x in model['nodes']};index=rl.load(model['inputs']['scientific_index']['path'])
    annotations=rl.load(annotations_path) if annotations_path else {}
    if annotations:
        if annotations.get('logic_snapshot')!=model['snapshot']:raise ValueError('WORKBOOK_ANNOTATION_DRIFT')
        for row in annotations.get('rows',[]):
            if not row.get('source_ids') or not set(row['source_ids'])<=set(model['sources']):raise ValueError('WORKBOOK_ANNOTATION_SOURCE')
    manifest=rl.load(manifest_path);graphs={x['view']:x for x in manifest['graphs'] if x['logic_snapshot']==model['snapshot']}
    required=['A','B']+[b['id'] for b in model['branches'][:2]]
    for name in required:
        rec=graphs[name]
        if rl.file_sha(rec['png'])!=rec['png_sha256'] or rl.file_sha(rec['pptx'])!=rec['pptx_sha256']:raise ValueError('WORKBOOK_GRAPH_DRIFT: '+name)
    output=Path(output)
    if output.exists():raise ValueError('OUTPUT_EXISTS')
    output.parent.mkdir(parents=True,exist_ok=True)
    doc=Document();page_setup(doc.sections[0]);doc.core_properties.title=model['paper_title']+' 科研逻辑拆解手册';doc.core_properties.identifier=model['snapshot'];doc.core_properties.subject='Shared research logic, Methods and evidence';doc.core_properties.author='Paper Figure Story'
    for sty in ['Normal','Caption','Heading 1','Heading 2','Title']:
        s=doc.styles[sty];s.font.name='Times New Roman';s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor.from_string('000000' if sty=='Title' else '174994' if sty.startswith('Heading') else '202A36');s.font.size=Pt(28 if sty=='Title' else 18 if sty=='Heading 1' else 13 if sty=='Heading 2' else 10.5)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.06
        s.font.bold=sty.startswith('Heading')
    footer=doc.sections[0].footer.paragraphs[0];fonts(footer.add_run('逻辑模型 '+model['snapshot'][:16]+'   '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    used=[]
    def para(txt,style=None):return doc.add_paragraph(txt,style)
    def heading(txt):doc.add_heading(txt,1)
    def newpage(txt):doc.add_page_break();heading(txt)
    def diagram(view,max_width=10.5,max_height=6.05):
        g=graphs[view];im=Image.open(g['png']);ratio=im.width/im.height;w=min(max_width,max_height*ratio);p=para('');p.paragraph_format.space_after=Pt(0);p.add_run().add_picture(g['png'],width=Inches(w));para('同一模型图 '+view+'   '+model['snapshot'][:16]+'   图形节点可在配套PPTX编辑','Caption');used.append({'view':view,'png':g['png'],'png_sha256':g['png_sha256'],'pptx':g['pptx'],'logic_snapshot':model['snapshot']})
    doc.add_heading(annotations.get('title','科研逻辑拆解手册' if inv else '科研论证拆解手册'),0)
    para(model['paper_title'])
    heading('01 核心问题与研究缺口')
    para(nm[model['core']['nodes'][0]]['short_text'])
    rows=[]
    for key in ['background_gap','core_question','hypotheses']:
        facet=index['facets'][key];sids=facet.get('sources',[])
        rows.append([{'background_gap':'已有认识与缺口','core_question':'中心问题','hypotheses':'待检验解释与边界'}[key],facet.get('text','见有源科学索引'),'p'+','.join(str(model['sources'][s]['physical_page']) for s in sids if s in model['sources'])])
    table(doc,['问题层次','科学内容','来源'],rows,[1.6,7.2,1.6],11)
    glossary=[r for r in annotations.get('rows',[]) if r.get('kind')=='glossary']
    if glossary:
        newpage('阅读术语')
        table(doc,['对象或术语','含义','来源'],[[r['term'],r['meaning'],'p'+','.join(str(model['sources'][s]['physical_page']) for s in r['source_ids'])] for r in glossary],[2.0,7.0,1.4],10)
    para(str(len(methods))+'条记录包含实验、控制、计算或共享任务，不等于同等数量的独立实验。图和表采用论文论证关系，不冒充实验执行日历。' if inv else '本手册整理综述论证与引用证据，不把引用研究登记成作者的新实验。')
    newpage('02 全文科研逻辑总览');diagram('A')
    # A larger page preserves readability of the independent detailed canvas.
    section=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(section,20,16);heading('完整科研逻辑与实验设计');diagram('B',18.3,13.0)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(section)
    for bi,branch in enumerate(model['branches'][:2]):
        if bi:doc.add_page_break()
        heading('03 '+branch['title']);diagram(branch['id'])
        newpage('实验设计 读数与结果对照')
        exids=list(dict.fromkeys(b['experiment_id'] for nid in branch['nodes'] for b in nm[nid]['field_bindings']))
        rows=[]
        for eid in exids:
            m=em[eid];p=m['plain_explanation'];pages=','.join(str(model['sources'][s]['physical_page']) for s in p['source_ids']);rows.append([eid+'\n'+m['name']+'\n'+','.join(m['figure_ids'])+' / p'+pages,p['why'],p['design']+'\n'+p['readout'],p['result'],p['meaning']+'\n'+p['limit']])
        table(doc,['任务与来源','为什么做','怎样比较和读图','实际发现','能说明和不能说明'],rows,[1.65,1.55,2.65,2.0,2.55],9.5)
    newpage('04 跨分支证据与综合认识')
    chosen=[e for e in model['edges'] if e['id'] in model['core']['edges']]
    for j in range(0,len(chosen),6):
        if j:doc.add_page_break();heading('跨分支关系与解释边界')
        rows=[]
        for e in chosen[j:j+6]:rows.append([e['id']+' '+nm[e['from']]['title']+' / '+nm[e['to']]['title'],e['type']+'\n'+e['identity'],e['rationale'],e['alternatives'],'p'+','.join(str(model['sources'][s]['physical_page']) for s in e['source_ids'])+'\nPARTIAL'])
        table(doc,['起点与终点','关系身份','判断理由','不能额外推出','来源与状态'],rows,[2.0,1.7,2.7,2.7,1.3],10)
    if inv:
        full=output.parent/'Full_Experimental_Inventory.source.json';shutil.copy2(model['inputs']['inventory']['path'],full)
        csvpath=output.parent/'Methods字段与来源完整索引.csv'
        with csvpath.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f);w.writerow(['Experiment ID','字段','完整值','字段状态','原始字段来源','原始限制'])
            for m in methods:
                for key,field in m['fields'].items():w.writerow([m['experiment_id'],key,value(field.get('value')),field.get('status'),json.dumps(field.get('source_refs',[]),ensure_ascii=False),value(field.get('reason'))])
        for j in range(0,len(methods),9):
            newpage('05 完整Methods结构化清单' if j==0 else 'Methods清单续表')
            rows=[]
            for m in methods[j:j+9]:
                f=m['fields'];p=m['plain_explanation'];counts='；'.join(k+'='+value(f.get(k,{}).get('value'))+' ['+f.get(k,{}).get('status','NOT_VERIFIED')+']' for k in ['sample_count','biological_replicates','technical_replicates','independent_runs'])
                rows.append([m['experiment_id']+'\n'+m['task_type'],m['name'],p['design']+'\n'+p['readout'],counts,f['methods']['status']+' / '+f['result']['status']+'\n'+','.join(m['figure_ids'])+'\n字段 '+str(len(f))+'项'])
            table(doc,['稳定ID与任务类别','任务名称','设计和关键检测','数量及重复的原记录','Methods/结果状态'],rows,[1.65,1.8,3.1,2.55,1.3],8.8)
        para('以上'+str(len(methods))+'项保留全部原始字段的结构化引用。完整材料、分组、处理、方法、分析、统计、Results及逐字段来源在配套文件中，不以压缩表替换原记录。')
        p=para('完整结构化Methods记录：');hyperlink(p,full.name,full.name)
        p=para('逐字段可编辑CSV：');hyperlink(p,csvpath.name,csvpath.name)
        para('原记录SHA256 '+rl.file_sha(full),'Caption')
    newpage('06 来源与核验状态索引')
    statusrows=[['文件 页码与摘录','已检查文件SHA及字符区间','定位成功不等于科学语义或独立终审通过'],['正文和Methods阅读',index['reading_status'],'Supplementary缺失不等于原文未报告'],['科学解释','PARTIAL；独立终审未进行','机制解释保留原身份，未知参数不填补']]
    statusrows += [[r['title'],r['detail'],r['impact']] for r in annotations.get('rows',[]) if r.get('kind')=='status']
    table(doc,['核验层次','当前状态','影响'],statusrows,[2.0,4.1,4.3],10.5)
    sources=[]
    for sid,s in model['sources'].items():sources.append([sid,str(s['physical_page']),Path(s['path']).name+' chars '+str(s['start'])+'–'+str(s['end']),s['source_status']+' / PARTIAL'])
    for j in range(0,len(sources),16):
        if j:doc.add_page_break();heading('来源索引续表')
        table(doc,['来源ID','物理页','缓存定位','状态'],sources[j:j+16],[2.0,.7,4.0,3.7],9.5)
    para('逻辑模型 '+model['snapshot']);para('配套PPTX中的节点、连线与本文图像来自同一模型版本。文档文字和表格可编辑；图像可替换，节点编辑在配套PPTX中完成。')
    # The bundled Word template may include paragraph rules; keep the title plain.
    for root in [doc.element,doc.styles.element]:
        for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.save(output)
    from apply_script_fonts import apply_docx_fonts
    apply_docx_fonts(output)
    receipt={'version':'pfs-workbook-1','docx':str(output),'docx_sha256':rl.file_sha(output),'logic_model':str(Path(model_path).resolve()),'logic_snapshot':model['snapshot'],'network_snapshot':model['network_snapshot'],'annotations_sha256':rl.file_sha(annotations_path) if annotations_path else None,'graphs':used,'inventory_source_sha256':rl.file_sha(full) if inv else None,'experiment_ids':[m['experiment_id'] for m in methods],'format':'Structured diagrams, comparisons and source tables; full Methods companion preserved, no article rewrite.','document_text_tables_editable':True,'diagram_editing':'Native nodes/arrows editable in linked PPTX; DOCX contains their exact exported PNGs.','render_status':'PENDING','formal_ready':False}
    output.with_suffix('.receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8');return receipt
def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--model',required=True,type=Path);p.add_argument('--graphs',required=True,type=Path);p.add_argument('--output',required=True,type=Path);p.add_argument('--annotations',type=Path);p.add_argument('--reader',action='store_true');p.add_argument('--sample',action='store_true');a=p.parse_args()
    if a.reader:
        from reader_workbook import build_reader
        r=build_reader(a.model,a.graphs,a.output,a.annotations,a.sample)
    else:r=build(a.model,a.graphs,a.output,a.annotations)
    print(json.dumps({'docx':r['docx'],'experiments':len(r['experiment_ids']),'graphs':len(r['graphs'])},ensure_ascii=False))
if __name__=='__main__':main()
