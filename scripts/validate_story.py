#!/usr/bin/env python3
"""Structural QA for Paper Figure Story outputs; never claims scientific verification."""
import argparse
import csv
import pathlib
import re
import sys

FILES = [
    '01_论文总体科研逻辑.md',
    '02_Figure逐图研究故事.md',
    '03_PPT逐页讲述脚本.md',
    '04_证据溯源与核查.md',
    '05_实验设计证据矩阵.csv',
]
REQUIRED_CSV = ['Figure', 'Panel', '问题来源', '科学问题', '样本与分组', '对照', '实验方法', '测量指标', '关键结果', '结论及证据级别', '下一步问题', '论文来源']


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('outdir', type=pathlib.Path)
    p.add_argument('--allow-partial', action='store_true', help='Validate only files present')
    p.add_argument('--profile', choices=['balanced','quick','full'], default='balanced', help='Output profile')
    args = p.parse_args()
    root = args.outdir
    errors, warns = [], []
    if args.profile in ('balanced','quick'):
        light_files = ['01_科研逻辑与逐Figure证据.md', '06_PPT页面数据.json', '07_科研逻辑汇报.pptx', '08_简明质检.md']
        for n in light_files:
            fp = root / n
            if not fp.exists():
                (warns if args.allow_partial else errors).append('缺少文件：' + n)
            elif not fp.stat().st_size:
                errors.append('文件为空：' + n)
        jpath = root / '06_PPT页面数据.json'
        if jpath.is_file():
            try:
                import json
                doc = json.loads(jpath.read_text(encoding='utf-8'))
                ss = doc.get('slides', [])
                if not isinstance(ss,list) or not ss:
                    errors.append('PPT JSON缺少有效slides')
                elif not any(x.get('kind')=='framework' for x in ss):
                    errors.append('PPT JSON缺少全文联合总框架(kind=framework)')
                if ss and not any(x.get('kind')=='figure' for x in ss):
                    errors.append('PPT JSON缺少逐Figure页')
            except Exception as e:
                errors.append('JSON解析失败：'+str(e))
        pptx = root / '07_科研逻辑汇报.pptx'
        if pptx.exists() and pptx.stat().st_size:
            try:
                from pptx import Presentation
                prs = Presentation(str(pptx))
                if len(prs.slides)<3:
                    errors.append('PPT页数过少：'+str(len(prs.slides)))
            except Exception as e:
                errors.append('PPT解析失败：'+str(e))
        print(f'检查目录：{root}  Profile: {args.profile}')
        for v in warns: print('WARN:',v)
        for v in errors: print('ERROR:',v)
        if not warns and not errors: print('PASS：文件与结构检查通过（不是科学内容核实）')
        return 1 if errors else 0
    for n in FILES:
        f = root / n
        if not f.exists():
            (warns if args.allow_partial else errors).append('缺少文件：' + n)
        elif not f.stat().st_size:
            errors.append('文件为空：' + n)
    fig = root / FILES[1]
    if fig.is_file():
        s = fig.read_text(encoding='utf-8')
        figures = re.findall(r'(?im)^\s*#{1,4}\s*(?:Figure|Fig\.?|图)\s*\d+', s)
        if not figures:
            warns.append('逐图文件没有可识别的 Figure 小节标题')
        for word in ['样本', '实验', '结果', '下一', '来源']:
            if word not in s:
                warns.append(f'逐图文件缺少关键词：{word}')
        if not re.search(r'\bFig(?:ure)?\.?\s*\d+\s*[a-z]?\b', s, re.I):
            warns.append('逐图报告缺少可核查的 Fig/panel 引用')
    ppt = root / FILES[2]
    if ppt.is_file():
        s = ppt.read_text(encoding='utf-8')
        if not re.search(r'(?im)^\s*#{1,4}\s*(?:Slide|幻灯片|第\s*\d+\s*页)', s):
            warns.append('PPT逐页文件没有可识别的页级标题')
        for w in ['科学问题', '技术路线', '实验', '结果']:
            if w not in s:
                warns.append(f'PPT逐页文件可能缺少：{w}')
    datafile = root / FILES[4]
    if datafile.is_file():
        try:
            with datafile.open('r', encoding='utf-8-sig', newline='') as f:
                reader = csv.DictReader(f)
                names = reader.fieldnames or []
                missing = [x for x in REQUIRED_CSV if x not in names]
                if missing:
                    errors.append('证据矩阵CSV缺少列：' + ', '.join(missing))
                rows = list(reader)
                if not rows:
                    warns.append('证据矩阵CSV没有数据行')
                else:
                    n_missing = sum(not r.get('论文来源', '').strip() for r in rows)
                    if n_missing:
                        warns.append(f'{n_missing}条证据缺少论文来源')
        except Exception as e:
            errors.append('读取CSV失败：' + str(e))
    print(f'检查目录：{root}')
    print('基础静态检查（不能代替原文事实核查）')
    for s in warns:
        print('WARN:', s)
    for s in errors:
        print('ERROR:', s)
    if not warns and not errors:
        print('PASS：结构检查通过；仍需核对原论文。')
    return 1 if errors else 0


if __name__ == '__main__':
    sys.exit(main())
