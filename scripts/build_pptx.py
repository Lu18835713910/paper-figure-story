#!/usr/bin/env python3
"""Render verified Figure-story JSON into editable PPTX; preserve tagged PPTX archetypes.

Usage:
  python build_pptx.py slides.json --output report.pptx
  python build_pptx.py slides.json --template custom.pptx --output report.pptx
  python build_pptx.py --make-default-template assets/default_academic_blue_174994.pptx

Requires python-pptx and Pillow. This script DOES NOT extract/certify facts from papers.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

try:
    from PIL import Image
except ImportError:
    Image = None

BASE = Path(__file__).resolve().parent.parent
DEFAULT_PATH = BASE / 'assets' / 'default_academic_blue_174994.pptx'
COLORS = {
    'accent': '174994', 'highlight': '3372C5', 'dark': '17345E',
    'light': 'EAF1FB', 'paper': 'F8FAFD', 'white': 'FFFFFF',
    'body': '40536B', 'border': 'D7E1F0',
}
CJK_FONT = 'Microsoft YaHei'
VALID_KINDS = {'cover', 'question', 'framework', 'logic', 'route', 'figure', 'conclusion', 'technology', 'text'}


def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def textbox(slide, name, x, y, w, h, text='', size=17, color='body', bold=False,
            align=PP_ALIGN.LEFT, fill=None, margin=.04, valign=MSO_ANCHOR.MIDDLE,
            radius=False, style=None):
    style = style or COLORS
    if fill:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                     Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(style.get(fill, fill))
        shp.line.fill.background()
        if radius:
            try: shp.adjustments[0] = 0.08
            except (IndexError, AttributeError): pass
    else:
        shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = name
    tf = shp.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(.015)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = CJK_FONT; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = rgb(style.get(color, color))
    return shp


def rectangle(slide, name, x, y, w, h, fill, style=None, rounded=False):
    style = style or COLORS
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name; s.fill.solid(); s.fill.fore_color.rgb = rgb(style.get(fill, fill))
    s.line.fill.background()
    if rounded:
        try: s.adjustments[0] = 0.07
        except (IndexError, AttributeError): pass
    return s


def add_frame(slide, kind, width, height, style=None):
    """Construct one editable style archetype, with named PFS text/image slots."""
    style = style or COLORS
    W, H = width, height
    sx, sy = W / 13.333, H / 7.5
    def t(name, x, y, w, h, **kw):
        return textbox(slide, name, x*sx, y*sy, w*sx, h*sy, style=style, **kw)
    def r(name, x, y, w, h, fill, **kw):
        return rectangle(slide, name, x*sx, y*sy, w*sx, h*sy, fill, style=style, **kw)

    if kind == 'cover':
        r('PFS_LAYOUT_cover', 0, 0, 13.333, 7.5, 'dark')
        r('DECOR_lowerband', 0, 6.63, 13.333, .87, 'accent')
        r('DECOR_accentline', 1.0, 1.58, .14, 3.65, 'highlight')
        r('DECOR_cornerline', 10.9, .88, 1.36, .07, 'highlight')
        t('PFS_SECTION', 1.22, 1.15, 10.8, .38, text='RESEARCH PRESENTATION',
          size=14, bold=True, color='highlight')
        t('PFS_TITLE', 1.22, 1.92, 10.65, 1.75, text='论文科研逻辑研究报告',
          size=35, color='white', bold=True)
        t('PFS_SUBTITLE', 1.24, 3.8, 10.45, 1.1, text='Paper title / Research topic',
          size=20, color='white')
        t('PFS_CITATION', 1.23, 6.8, 9.7, .35, text='Author · Journal · Year',
          size=12, color='white')
        return

    r('PFS_LAYOUT_'+kind, 0, 0, 13.333, 7.5, 'paper')
    r('DECOR_top', 0, 0, 13.333, .15, 'accent')
    r('DECOR_titleline', .7, 1.52, 11.95, .027, 'border')
    t('PFS_SECTION', .76, .28, 6.0, .28, text='RESEARCH LOGIC / FIGURE STORY',
      size=10, color='accent', bold=True)
    t('PFS_TITLE', .72, .67, 11.93, .75, text='科学问题与实验结果',
      size=27, bold=True, color='dark')
    r('DECOR_footerline', .72, 7.13, 11.9, .018, 'border')
    t('PFS_CITATION', .76, 7.19, 9.7, .22, text='Author · Journal · Year',
      color='body', size=9)
    t('PFS_NUMBER', 11.58, 7.18, .82, .24, text='01', size=9,
      color='accent', align=PP_ALIGN.RIGHT, bold=True)

    if kind == 'question':
        t('PFS_LEAD', .8, 1.76, 11.4, 1.05, text='研究领域已经知道什么？现有证据还不能回答什么？',
          size=19, color='body')
        r('DECOR_card', .72, 3.06, 11.88, 3.75, 'white', rounded=True)
        r('DECOR_cardbar', .72, 3.12, .065, 3.57, 'accent')
        t('PFS_BULLETS', 1.1, 3.36, 11.0, 2.95,
          text='科学问题一\n科学问题二\n科学问题三', size=21, color='dark')
    elif kind == 'framework':
        t('PFS_LEAD', .78, 1.67, 11.60, .49, text='中心科学问题：____ ｜ 研究对象与处理：____', size=14, color='body')
        r('PFS_GRAPH_AREA', .73, 2.22, 11.85, 4.52, 'white', rounded=True)
        t('PFS_TRANSITION', .84, 6.81, 11.65, .22, text='', size=10, color='body')
    elif kind in ('logic', 'route'):
        t('PFS_LEAD', .78, 1.75, 11.6, .56, text='依照论文证据关系生成研究逻辑图',
          size=17, color='body')
        r('PFS_GRAPH_AREA', .74, 2.47, 11.87, 3.88, 'white', rounded=True)
        t('PFS_TRANSITION', .85, 6.48, 11.65, .41,
          text='注：分支、并行验证及机制推断应依照论文实际证据设置。', size=12, color='body')
    elif kind == 'figure':
        r('DECOR_figurebg', .7, 1.72, 7.00, 5.19, 'white', rounded=True)
        r('PFS_IMAGE', .85, 1.86, 6.69, 4.65, 'white')
        t('PFS_IMAGE_PLACEHOLDER', 1.1, 3.53, 6.1, .8,
          text='原论文 Figure / panel 图像区域', size=17,
          color='accent', align=PP_ALIGN.CENTER)
        t('PFS_FIGURE_REF', .95, 6.56, 6.48, .23,
          text='Fig 2a–c · Original paper p.5', size=10, color='body')
        r('DECOR_rightbg', 7.83, 1.72, 4.77, 5.19, 'white', rounded=True)
        t('DECOR_q_label', 8.87, 1.91, 3.4, .27,
          text='研究问题', size=11, color='accent', bold=True)
        t('PFS_QUESTION', 8.87, 2.19, 3.47, .73,
          text='这个实验想解决什么问题？', size=16, color='dark')
        r('DECOR_div_a', 8.88, 3.01, 3.36, .018, 'border')
        t('DECOR_d_label', 8.87, 3.13, 3.4, .26,
          text='样本 / 方法 / 对照', size=11, color='accent', bold=True)
        t('PFS_DESIGN', 8.87, 3.44, 3.44, .84,
          text='样本与实验设计、主要读数', size=14, color='body')
        r('DECOR_div_b', 8.88, 4.36, 3.36, .018, 'border')
        t('DECOR_r_label', 8.87, 4.49, 3.4, .27,
          text='关键结果 / 证据强度', size=11, color='accent', bold=True)
        t('PFS_RESULT', 8.87, 4.78, 3.44, .94,
          text='图中看到什么；证据支持到哪一步', size=14, color='dark')
        r('DECOR_div_c', 8.88, 5.8, 3.36, .018, 'border')
        t('PFS_TRANSITION', 8.87, 5.94, 3.47, .66,
          text='下一步：仍需回答……', size=12, color='accent')
    elif kind == 'conclusion':
        r('DECOR_conclusion_bg', .75, 1.78, 7.55, 4.95, 'white', rounded=True)
        r('DECOR_conclusion_right', 8.55, 1.78, 4.0, 4.95, 'light', rounded=True)
        t('DECOR_findings', 1.06, 2.11, 6.9, .36,
          text='论文支持的主要发现', size=15, color='accent', bold=True)
        t('PFS_POINTS', 1.07, 2.75, 6.73, 3.55,
          text='发现一\n发现二\n发现三', size=18, color='dark')
        t('DECOR_limit_label', 8.87, 2.11, 3.2, .35,
          text='证据边界与局限', size=15, color='accent', bold=True)
        t('PFS_LIMITATION', 8.87, 2.69, 3.3, 3.25,
          text='哪些是统计关联？哪些仍是机制假说？', size=16, color='dark')
    elif kind in ('technology', 'text'):
        r('DECOR_bodybg', .76, 1.8, 11.77, 4.95, 'white', rounded=True)
        t('PFS_LEAD', 1.05, 2.0, 10.95, .58,
          text='这一技术在本研究中提供了什么关键观测量？', size=19, color='accent')
        t('PFS_POINTS', 1.08, 2.82, 10.78, 3.35,
          text='技术价值一\n技术价值二\n技术价值三', size=18, color='dark')
    else:
        raise ValueError('Unsupported kind: '+kind)


def make_default_template(path):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = next((lay for lay in prs.slide_layouts if 'blank' in lay.name.lower()), prs.slide_layouts[-1])
    for kind in ['cover', 'question', 'framework', 'logic', 'route', 'figure', 'conclusion', 'technology', 'text']:
        slide = prs.slides.add_slide(blank)
        add_frame(slide, kind, 13.333, 7.5)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def remove_slide(prs, slide_id):
    sldIdLst = prs.slides._sldIdLst
    rel_id = slide_id.rId
    sldIdLst.remove(slide_id)
    prs.part.drop_rel(rel_id)


def duplicate_slide(prs, src):
    dest = prs.slides.add_slide(src.slide_layout)
    # Eliminate any layout-added placeholders; clone actual shapes and all image relationships.
    for shp in list(dest.shapes):
        dest.shapes._spTree.remove(shp.element)
    for shp in src.shapes:
        el = copy.deepcopy(shp.element)
        for elem in el.iter():
            for key, val in list(elem.attrib.items()):
                if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    try:
                        rel = src.part.rels[val]
                        target = rel.target_ref if rel.is_external else rel.target_part
                        elem.set(key, dest.part.relate_to(target, rel.reltype, is_external=rel.is_external))
                    except (KeyError, AttributeError):
                        pass
        dest.shapes._spTree.insert_element_before(el, 'p:extLst')
    # Copy per-slide background/other cSld non-shape data if present.
    src_bg = src._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if src_bg is not None:
        dst_bg = dest._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
        if dst_bg is not None: dest._element.cSld.remove(dst_bg)
        dest._element.cSld.insert(0, copy.deepcopy(src_bg))
    return dest


def find_archetypes(prs):
    idx = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name.startswith('PFS_LAYOUT_'):
                kind = shape.name[len('PFS_LAYOUT_'):]
                if kind in VALID_KINDS and kind not in idx:
                    idx[kind] = slide
    return idx


def theme_rgb_from_shape(prs):
    """Best-effort sample from untagged PPTX; not a promise of exact slide designs."""
    st = dict(COLORS)
    candidate = []
    fonts = []
    def legible_on_white(value):
        # Decorative pale rules must not become the primary text/arrow color.
        channels = [int(value.lstrip('#')[i:i+2], 16) / 255 for i in (0, 2, 4)]
        linear = [c / 12.92 if c <= .04045 else ((c + .055) / 1.055) ** 2.4 for c in channels]
        luminance = sum(c * w for c, w in zip(linear, (.2126, .7152, .0722)))
        return 1.05 / (luminance + .05) >= 4.5
    for slide in list(prs.slides)[:min(5, len(prs.slides))]:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                try:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            col = run.font.color
                            if col and col.type and col.rgb:
                                if shape.top < Inches(1.8) and run.font.size and run.font.size >= Pt(16):
                                    fonts.append('#'+str(col.rgb))
                except (ValueError, AttributeError, TypeError):
                    pass
            if shape.shape_type != 1:
                continue
            try:
                col = shape.fill.fore_color
                if shape.fill.type and col.type and col.rgb:
                    if shape.height < Inches(.45) and shape.width > Inches(2.0):
                        candidate.append('#'+str(col.rgb))
            except (ValueError, AttributeError, TypeError):
                continue
    candidate = [c for c in candidate if legible_on_white(c)]
    fonts = [c for c in fonts if legible_on_white(c)]
    if candidate: st['accent'] = Counter(candidate).most_common(1)[0][0]
    if fonts: st['dark'] = Counter(fonts).most_common(1)[0][0]
    return st


def clear_shape_fill_text(shape):
    if not shape.has_text_frame: return
    shape.text_frame.clear()


def styled_text(shape, value, is_list=False):
    if not shape.has_text_frame: return
    # Preserve the named template slot's first-run style rather than impose default color/font.
    fsize, fcolor, fname, fbold, align = None, None, CJK_FONT, False, PP_ALIGN.LEFT
    for p in shape.text_frame.paragraphs:
        if p.runs:
            f = p.runs[0].font
            fsize = f.size; fname = f.name or CJK_FONT; fbold = bool(f.bold)
            try:
                if f.color and f.color.type and f.color.rgb: fcolor = f.color.rgb
            except (ValueError, AttributeError): pass
            align = p.alignment or PP_ALIGN.LEFT
            break
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    if isinstance(value, list): lines = [str(x) for x in value if str(x).strip()]
    else: lines = [str(value or '')]
    if not lines: lines = ['']
    basept = float(fsize.pt if fsize else 16)
    n_chars = sum(max(1, len(x)) for x in lines)
    width_pt = max(1, shape.width / 12700)  # inches*72 pt; 914400 emu/in = 12700 emu/pt
    height_pt = max(1, shape.height / 12700)
    # Fast conservative approximate fit for Chinese text; preserve size when possible.
    available = width_pt * height_pt
    estimate_pt = (available / max(1, n_chars) / 1.5)**.5
    fontsize = max(9, min(basept, estimate_pt*1.1)) if n_chars>18 else basept
    for i, s in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = Pt(0)
        p.space_after = Pt(8 if is_list else 1)
        p.level = 0
        run = p.add_run(); run.text = ('•  ' + s) if is_list else s
        run.font.name = fname; run.font.bold = fbold
        run.font.size = Pt(fontsize)
        if fcolor is not None: run.font.color.rgb = fcolor
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if shape.name in ('PFS_TITLE','PFS_SECTION') else MSO_ANCHOR.TOP


def insert_images(slide, shape, paths, json_path, allow_missing):
    if len(paths)>2:
        raise ValueError('TOO_MANY_IMAGES: 此成熟版式每页最多2幅原图；请按科学问题拆页，不能静默丢弃后续原图。')
    blank_note = next((v for v in slide.shapes if v.name == 'PFS_IMAGE_PLACEHOLDER'), None)
    if blank_note is not None:
        clear_shape_fill_text(blank_note)
    if not paths:
        if not allow_missing:
            raise ValueError('Figure页缺少原图image/images；正式成品必须提供Figure面板图片')
        if blank_note is not None: styled_text(blank_note, 'FIGURE IMAGE MISSING — DRAFT')
        return
    resolved = []
    for p in paths:
        fp = Path(p)
        if not fp.is_absolute(): fp = json_path.parent / fp
        if not fp.is_file():
            if not allow_missing: raise FileNotFoundError(f'找不到Figure图片：{fp}')
            continue
        resolved.append(fp)
    if not resolved:
        if blank_note is not None: styled_text(blank_note, 'FIGURE IMAGE MISSING — DRAFT')
        return
    # Add images without crop, preserve aspect and all panel coordinates/statistics.
    x, y, width, height = shape.left, shape.top, shape.width, shape.height
    n = len(resolved)
    for i, fp in enumerate(resolved):
        cx = x + i*(width//n)
        cw = width//n - (Inches(.06) if n > 1 else 0)
        with Image.open(fp) as im:
            imw, imh = im.size
        scale = min(cw / imw, height / imh)
        aw, ah = int(imw*scale), int(imh*scale)
        lx = cx + (cw-aw)//2
        ty = y + (height-ah)//2
        slide.shapes.add_picture(str(fp), lx, ty, width=aw, height=ah)



# Figure-grounded, editable scientific argument graphs.
# Nodes are scientific questions, experimental designs and observations;
# edges explicitly encode author-supported vs presentation-inferred relationships.
EDGE_TYPES = {'motivates', 'tests', 'supports', 'validates', 'compares',
              'parallel', 'extends', 'feedback', 'integrates', 'alternative'}
NODE_ROLES = {'background', 'question', 'experiment', 'observation',
              'analysis', 'validation', 'inference', 'conclusion'}


def _graph_bounds(slide):
    area=next((sh for sh in slide.shapes if sh.name == 'PFS_GRAPH_AREA'),None)
    if area is not None:
        return tuple(v/914400 for v in (area.left,area.top,area.width,area.height))
    legacy=[sh for sh in slide.shapes if sh.name.startswith('PFS_STAGE_') or sh.name.startswith('DECOR_stage')]
    if legacy:
        x=min(v.left for v in legacy)/914400
        y=min(v.top for v in legacy)/914400
        x2=max(v.left+v.width for v in legacy)/914400
        y2=max(v.top+v.height for v in legacy)/914400
        return (x, y, x2-x, y2-y)
    raise ValueError('自定义模板的logic/route页面需要PFS_GRAPH_AREA插槽；旧版PFS_STAGE_*布局可自动升级。')


def _graph_layout(nodes, edges, W, H):
    # Explicit normalized card-center coordinates preserve purposeful user/story arrangement.
    if all(isinstance(n.get('x'), (int,float)) and isinstance(n.get('y'), (int,float)) for n in nodes):
        for n in nodes:
            if not (0 <= n['x'] <= 1 and 0 <= n['y'] <= 1):
                raise ValueError(f'逻辑图位置超出0..1: {n["id"]}')
        cols = max(2, len({round(n['x'],2) for n in nodes}))
        max_rows = max(1, len({round(n['y'],2) for n in nodes}))
        pos={n['id']:(n['x'],n['y']) for n in nodes}
    else:
        # Longest-path topological layers; parallel and convergent relationships are not forced linear.
        indeg={n['id']:0 for n in nodes}
        forward={n['id']:[] for n in nodes}
        for e in edges:
            if e.get('relation')=='feedback': continue
            forward[e['from']].append(e['to']);indeg[e['to']]+=1
        queue=[n['id'] for n in nodes if indeg[n['id']]==0]
        order=[];layers={k:0 for k in indeg}
        while queue:
            cur=queue.pop(0);order.append(cur)
            for nxt in forward[cur]:
                layers[nxt]=max(layers[nxt],layers[cur]+1)
                indeg[nxt]-=1
                if indeg[nxt]==0:queue.append(nxt)
        if len(order)!=len(nodes):
            raise ValueError('逻辑图存在非feedback循环边；请修正方向或给循环边relation=feedback')
        cols=max(layers.values())+1
        group={k:[] for k in range(cols)}
        for n in nodes:group[layers[n['id']]].append(n['id'])
        max_rows=max(map(len,group.values()))
        pos={}
        for col,ids in group.items():
            for row,nid in enumerate(ids):
                pos[nid]=(0.5 if cols==1 else col/(cols-1), (row+1)/(len(ids)+1))
    if cols>5 or max_rows>3 or len(nodes)>10:
        raise ValueError('逻辑图过密：每页建议至多5列、3行、10节点。请依据子问题拆成两页，不要压缩关键科研信息。')
    cw = min(2.55, (W-.30-(cols-1)*.22)/max(1,cols))
    ch = min(1.45, (H-.27-(max_rows-1)*.17)/max(1,max_rows))
    if cw<1.65 or ch<.88:
        raise ValueError('逻辑图节点过多导致字体不可读，请拆页')
    return pos,cw,ch


def _connect(slide,x1,y1,x2,y2,style,dashed=False,arrow=False):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1), Inches(x2),Inches(y2))
    sh.name='PFS_GRAPH_EDGE'
    sh.line.color.rgb=rgb(style['border'] if dashed else style['accent'])
    sh.line.width=Pt(1.3 if dashed else 1.8)
    if dashed: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if arrow:
        from pptx.oxml.ns import qn
        ln=sh._element.spPr.find(qn('a:ln'))
        if ln is not None:
            end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');ln.append(end)
    return sh


def render_argument_graph(slide, record, style=None):
    style=style or COLORS
    nodes=record.get('nodes',[]);edges=record.get('edges',[])
    if not nodes:raise ValueError('研究思路页必须是有根据的科研关系图，请提供nodes/edges，不得只填线性stages。')
    ids={n['id'] for n in nodes}
    if len(ids)!=len(nodes):raise ValueError('逻辑图节点id重复')
    x0,y0,W,H=_graph_bounds(slide)
    # Remove only the old fixed five-stage content when retrofitting legacy tagged templates.
    for shp in list(slide.shapes):
        if shp.name.startswith('PFS_STAGE_') or shp.name.startswith('DECOR_stage') or shp.name.startswith('DECOR_arrow'):
            slide.shapes._spTree.remove(shp.element)
    pos,cw,ch=_graph_layout(nodes,edges,W,H)
    coords={}
    for n in nodes:
        cx=x0+cw/2+.14+(W-cw-.28)*pos[n['id']][0]
        cy=y0+ch/2+.12+(H-ch-.24)*pos[n['id']][1]
        coords[n['id']]=(cx,cy)
    # Reject overlapping node geometry rather than exporting an unreadable paper logic slide.
    positions=list(coords.items())
    for i,(id_a,(ax,ay)) in enumerate(positions):
        for id_b,(bx,by) in positions[i+1:]:
            if abs(ax-bx)<cw*.90 and abs(ay-by)<ch*.90:
                raise ValueError(f'科研逻辑节点发生重叠: {id_a}/{id_b}；请调整x/y或拆成多页')
    # Edges beneath cards. Solid = paper-grounded, dashed = presentation inference.
    for e in edges:
        src=coords[e['from']];dst=coords[e['to']]
        inferred=e.get('provenance')=='presentation'
        feedback=e.get('relation')=='feedback' or dst[0]<=src[0]-.05
        direction=1 if dst[0]>src[0] else -1
        xa=src[0]+direction*cw*.50; xb=dst[0]-direction*cw*.50
        ya,yb=src[1],dst[1]
        if feedback:
            # Feedback loop is a returning hypothesis check; indicate with dashed line.
            inferred=True
            yt=y0+H-.12
            _connect(slide,src[0],ya+ch*.5,src[0],yt,style,True)
            _connect(slide,src[0],yt,dst[0],yt,style,True)
            _connect(slide,dst[0],yt,dst[0],yb+ch*.5,style,True,True)
        elif abs(ya-yb)<.09:
            _connect(slide,xa,ya,xb,yb,style,inferred,True)
        else:
            mid=(xa+xb)/2
            _connect(slide,xa,ya,mid,ya,style,inferred)
            _connect(slide,mid,ya,mid,yb,style,inferred)
            _connect(slide,mid,yb,xb,yb,style,inferred,True)
        # Descriptive links live in their own small labels, not as invented causal arrows.
        lbl=str(e.get('label','')).strip()
        if lbl and abs(xb-xa)>.33:
            tx=(xa+xb)/2-.48; ty=(ya+yb)/2-.14
            if e.get('label_y') is not None:ty=y0+H*float(e['label_y'])
            textbox(slide,'PFS_EDGE_LABEL',max(x0+.03,min(x0+W-.98,tx)),ty,.95,.26,
                    lbl[:15],size=8.5,color='body',fill='white',align=PP_ALIGN.CENTER,
                    style=style,margin=.01)
    for n in nodes:
        cx,cy=coords[n['id']]
        x,y=cx-cw/2,cy-ch/2
        role=n.get('role','analysis')
        fill='light' if role in ('question','inference') else 'white'
        rectangle(slide,'PFS_GRAPH_NODE_'+n['id'],x,y,cw,ch,fill,style=style,rounded=True)
        rectangle(slide,'PFS_GRAPH_NODE_BAR_'+n['id'],x,y,cw,.072,
                  'highlight' if role in ('validation','conclusion') else 'accent',style=style)
        title=str(n.get('title') or n.get('label') or '')
        detail=str(n.get('detail','')).strip()
        fig=str(n.get('figure_ref','')).strip()
        if not title or not detail:
            raise ValueError(f'逻辑图节点{n["id"]}必须有科学问题/发现标题和具体实验/结论说明')
        textbox(slide,'PFS_GRAPH_NODE_TITLE',x+.11,y+.13,cw-.22,min(.57,ch*.38),title,
                size=12.5 if cw>2 else 11,color='dark',bold=True,style=style)
        textbox(slide,'PFS_GRAPH_NODE_DETAIL',x+.11,y+.55,cw-.22,max(.31,ch-.79),detail,
                size=10.2 if cw>2 else 9.3,color='body',style=style,valign=MSO_ANCHOR.TOP)
        if fig:
            textbox(slide,'PFS_GRAPH_FIGREF',x+.1,y+ch-.22,cw-.2,.18,fig,
                    size=8,color='accent',style=style)


def validate_framework(record, slide_i):
    """One research unit MUST integrate question, concrete design, and finding."""
    rows = record.get('research_units')
    if not isinstance(rows, list) or not 2 <= len(rows) <= 4:
        raise ValueError(f'第{slide_i}页综合研究框架要求2–4条research_units；过多请拆解为后续详图。')
    ids = []
    for r in rows:
        if not isinstance(r, dict) or any(not str(r.get(k, '')).strip() for k in ('id', 'question', 'design', 'finding', 'figure_ref')):
            raise ValueError(f'第{slide_i}页每条研究任务必须有id/question/design/finding/figure_ref，不能仅提供技术名。')
        ids.append(r['id'])
    if len(set(ids)) != len(ids): raise ValueError('研究任务id重复')
    for link in record.get('links', []):
        src = link.get('from', [])
        if isinstance(src, str): src = [src]
        if not src or any(item not in ids for item in src) or link.get('to') not in ids:
            raise ValueError(f'第{slide_i}页研究任务链接不存在: {link}')
        if link.get('relation') not in EDGE_TYPES:
            raise ValueError(f'研究任务关系必须是{sorted(EDGE_TYPES)}')
        if link.get('provenance') not in ('paper','presentation'):
            raise ValueError('链接必须标明是原文依据paper还是汇报组织presentation')


def render_integrated_framework(slide, record, style=None):
    """Author-grounded combined QUESTION -> DESIGN -> FINDING board.

    Each editable research row binds why/how/what. Cross-row relationships are
    recorded and captioned separately so parallel experiments are not passed off
    as a chronological chain.
    """
    style = style or COLORS
    rows = record['research_units']
    x0,y0,W,H = _graph_bounds(slide)
    N = len(rows)
    pad=.12
    head_h=.35
    gap=.095 if N>=4 else .16
    row_h=(H-head_h-2*pad-(N-1)*gap)/N
    widths=[W*.255,W*.365,W*.315]
    gaps=[W*.032,W*.024]
    sx=x0+pad
    c1=sx
    c2=c1+widths[0]+gaps[0]
    c3=c2+widths[1]+gaps[1]
    # Column headings convey the epistemic stages, not technology-only slots.
    for name,x,w in [('为什么研究（科学问题）',c1,widths[0]),('怎么验证（样本·分组·实验）',c2,widths[1]),('获得什么证据（结果·Figure）',c3,widths[2])]:
        textbox(slide,'PFS_FRAMEWORK_HEADER',x,y0+.055,w,.27,name,size=11.3,color='accent',bold=True,style=style)
    row_positions={}
    for i,r in enumerate(rows):
        y=y0+head_h+pad+i*(row_h+gap)
        row_positions[r['id']]=(y,row_h,i)
        prefix=f'{r["id"]}  '
        base='light' if i%2==0 else 'paper'
        rectangle(slide,'PFS_RESEARCH_QUESTION_'+r['id'],c1,y,widths[0],row_h,base,style=style,rounded=True)
        rectangle(slide,'PFS_RESEARCH_DESIGN_'+r['id'],c2,y,widths[1],row_h,'white',style=style,rounded=True)
        rectangle(slide,'PFS_RESEARCH_FINDING_'+r['id'],c3,y,widths[2],row_h,base,style=style,rounded=True)
        q=str(r['question']).strip()
        d=str(r['design']).strip()
        f=str(r['finding']).strip()
        if any(len(v)>95 for v in (q,d,f)):
            raise ValueError(f'研究任务{r["id"]}单栏文字过多，拆分或缩写：单栏建议不超过95字符。')
        textbox(slide,'PFS_RESEARCH_QUESTION_TEXT_'+r['id'],c1+.10,y+.07,widths[0]-.18,row_h-.13,prefix+q,
                size=11.4 if N<4 else 10.6,color='dark',bold=True,style=style,valign=MSO_ANCHOR.TOP)
        textbox(slide,'PFS_RESEARCH_DESIGN_TEXT_'+r['id'],c2+.10,y+.07,widths[1]-.18,row_h-.12,d,
                size=11.0 if N<4 else 10.4,color='body',style=style,valign=MSO_ANCHOR.TOP)
        textbox(slide,'PFS_RESEARCH_FINDING_TEXT_'+r['id'],c3+.10,y+.06,widths[2]-.18,row_h-.30,f,
                size=11.0 if N<4 else 10.4,color='dark',style=style,valign=MSO_ANCHOR.TOP)
        textbox(slide,'PFS_RESEARCH_FIGURE_'+r['id'],c3+.10,y+row_h-.21,widths[2]-.18,.18,
                str(r['figure_ref']),size=8.7,color='accent',style=style)
        _connect(slide,c1+widths[0]+.02,y+row_h/2,c2-.025,y+row_h/2,style,arrow=True)
        _connect(slide,c2+widths[1]+.02,y+row_h/2,c3-.025,y+row_h/2,style,arrow=True)
    # Adjacent supported question-to-question turns are visible as small editable arrows.
    # Do NOT connect parallel, distant, or unsupported research modules forcibly.
    for link in record.get('links',[]):
        prev=link.get('from',[])
        prev=[prev] if isinstance(prev,str) else prev
        if len(prev)!=1 or prev[0] not in row_positions or link.get('to') not in row_positions:
            continue
        sy,sh,si=row_positions[prev[0]]; ty,th,ti=row_positions[link['to']]
        if ti!=si+1 or link.get('relation') in ('parallel','compares','alternative'):
            continue
        cx=c1+widths[0]/2
        if link.get('provenance')=='paper':
            arr=slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(cx-.065), Inches(sy+sh+.015), Inches(.13), Inches(max(.055,gap-.03)))
            arr.name='PFS_FRAMEWORK_SUPPORTED_ARROW_'+prev[0]+'_'+link['to']
            arr.fill.solid();arr.fill.fore_color.rgb=rgb(style['accent']);arr.line.fill.background()
        else:
            seg=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(cx),Inches(sy+sh+.01),Inches(cx),Inches(ty-.01))
            seg.name='PFS_FRAMEWORK_PRESENTATION_LINK_'+prev[0]+'_'+link['to']
            seg.line.color.rgb=rgb(style['accent']);seg.line.width=Pt(1.2)
            seg.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    # Make cross-row synthesis explicit; never fake a straight experimental chronology.
    links=[]
    for lk in record.get('links',[]):
        src=lk['from']; src=[src] if isinstance(src,str) else src
        symbol = '+' if lk['relation']=='integrates' else ' / '
        label = lk.get('label') or lk['relation']
        marker='（汇报整理）' if lk.get('provenance')=='presentation' else ''
        links.append(symbol.join(src)+'→'+str(lk['to'])+':'+str(label)+marker)
    relationship='；'.join(links)[:145]
    if record.get('synthesis'): relationship=relationship+' ｜ 综合：'+str(record['synthesis'])[:45]
    foot=next((sh for sh in slide.shapes if sh.name=='PFS_TRANSITION'),None)
    if foot is not None: styled_text(foot,relationship or str(record.get('transition','各研究支线以原论文支持的关系组织。')))


# Paper-level MASTER research framework: study subject + treatment groups + parallel
# methods and their research functions + independent validation + final synthesis.
# This is NOT the per-Figure narrative chain. It mirrors users' hand-authored
# overview pages and remains fully editable in PowerPoint.
def validate_study_map(record, slide_i):
    m=record.get('study_map')
    if not isinstance(m,dict):
        raise ValueError(f'第{slide_i}页必须提供study_map对象')
    for field in ('material','groups','core_assays','validations','synthesis'):
        if not m.get(field):
            raise ValueError(f'第{slide_i}页总框架缺少study_map.{field}')
    if not isinstance(m['material'],dict) or not m['material'].get('name'):
        raise ValueError('study_map.material需要研究对象name')
    if not 2<=len(m['groups'])<=4:raise ValueError('总体框架分组建议2–4组；复杂设计应另加细化页')
    if not 2<=len(m['core_assays'])<=4:raise ValueError('总体框架核心实验应有2–4项')
    if not 2<=len(m['validations'])<=5:raise ValueError('总体框架验证实验应有2–5项')
    for a in m['core_assays']+m['validations']:
        if not isinstance(a,dict) or not a.get('name') or not a.get('purpose'):
            raise ValueError('每个技术项必须同时包含name（方法）和purpose（针对什么问题），仅列技术名称不允许')
    for g in m['groups']:
        if not g.get('name') or not g.get('detail'):
            raise ValueError('每个研究分组需name和detail（时间/处理/对照）')
    if len(m['synthesis'])>130:raise ValueError('总机制过长，应拆成结论要点/另页，不可压缩到小字号')


def render_study_map(slide,record,style=None):
    """Modular all-paper overview with editable figure/method/evidence blocks."""
    style=style or COLORS
    m=record['study_map']; x0,y0,W,H=_graph_bounds(slide)
    # Use fractions of graph area so 16:9 user template dimensions work.
    u=W/11.85; v=H/4.52
    def X(z):return x0+z*u
    def Y(z):return y0+z*v
    def box(name,x,y,w,h,fill='white',rounded=True):
        return rectangle(slide,name,X(x),Y(y),w*u,h*v,fill,style=style,rounded=rounded)
    def t(name,x,y,w,h,txt,size=11,color='dark',bold=False,align=PP_ALIGN.LEFT):
        return textbox(slide,name,X(x),Y(y),w*u,h*v,str(txt),size=size,color=color,bold=bold,
                       align=align,style=style,margin=.06,valign=MSO_ANCHOR.MIDDLE)
    def line(name,xa,ya,xb,yb,dash=False,arrow=False):
        z=_connect(slide,X(xa),Y(ya),X(xb),Y(yb),style,dash,arrow)
        z.name=name
        return z
    # Macro layout: material | treatment | core methods | validation.
    cols=[(.19,1.64),(2.12,2.16),(4.59,4.10),(9.02,2.61)]
    for idx,(a,w) in enumerate(cols):
        t('PFS_MASTER_COLUMN_LABEL_'+str(idx),a,.09,w,.28,
          ['研究对象 / 关键因子','处理与样本设计','核心技术及其科研任务','进一步验证'][idx],size=10.5,color='accent',bold=True,
          align=PP_ALIGN.CENTER)
    # Material/targets. Do not turn genes into a treatment or assert a physical mechanism.
    a,w=cols[0]
    box('PFS_MASTER_MATERIAL',a,.49,w,2.49,'light')
    t('PFS_MASTER_MATERIAL_NAME',a+.10,.85,w-.20,.53,m['material']['name'],size=23,bold=True,align=PP_ALIGN.CENTER)
    t('PFS_MASTER_MATERIAL_DETAIL',a+.10,1.42,w-.20,.53,m['material'].get('detail',''),size=11.1,color='body',align=PP_ALIGN.CENTER)
    if m['material'].get('targets'):
        box('PFS_MASTER_TARGETS_BG',a+.10,2.13,w-.20,.66,'white')
        t('PFS_MASTER_TARGETS',a+.14,2.16,w-.28,.59,
           '关注：'+str(m['material']['targets']),size=10.5,align=PP_ALIGN.CENTER)
    # Grouping card: clinical/tissue studies can instead use cohort/site as groups.
    a,w=cols[1]
    box('PFS_MASTER_TREATMENT',a,.49,w,2.49,'white')
    groups=m['groups']; ngh=len(groups)
    gh=min(.84, 2.02/ngh)
    for i,g in enumerate(groups):
        y=.60+i*(2.20/ngh)
        box('PFS_MASTER_GROUP_'+str(i),a+.12,y,w-.24,gh,'light' if i%2==0 else 'paper')
        t('PFS_MASTER_GROUP_NAME_'+str(i),a+.21,y+.045,w-.42,.26,g['name'],size=11.6,bold=True,align=PP_ALIGN.CENTER)
        t('PFS_MASTER_GROUP_DETAIL_'+str(i),a+.21,y+.30,w-.42,max(.28,gh-.34),g['detail'],size=10.4,color='body',align=PP_ALIGN.CENTER)
    # Parallel methods + purpose: these methods are not sequential stages.
    a,w=cols[2]
    box('PFS_MASTER_ASSAY_TABLE',a,.49,w,2.49,'white')
    assays=m['core_assays']; aheight=2.47/len(assays)
    for i,z in enumerate(assays):
        y=.51+i*aheight
        box('PFS_MASTER_ASSAY_NAME_BG_'+str(i),a+.035,y,w*.37,aheight-.045,'light')
        t('PFS_MASTER_ASSAY_NAME_'+str(i),a+.07,y+.07,w*.32,aheight-.16,z['name'],
          size=11.3,bold=True,align=PP_ALIGN.CENTER)
        t('PFS_MASTER_ASSAY_PURPOSE_'+str(i),a+w*.39,y+.07,w*.58,aheight-.16,
          z['purpose'],size=11.0,align=PP_ALIGN.CENTER)
        if i<len(assays)-1:
            box('PFS_MASTER_ASSAY_DIVIDER_'+str(i),a+.04,y+aheight-.026,w-.08,.012,'border',rounded=False)
    # Verification deliberately separate from primary multiomics readout.
    a,w=cols[3]
    box('PFS_MASTER_VALIDATION',a,.49,w,2.49,'light')
    vals=m['validations']; vh=2.28/len(vals)
    for i,z in enumerate(vals):
        y=.55+i*vh
        t('PFS_MASTER_VALIDATION_NAME_'+str(i),a+.08,y,w*.37,vh-.055,z['name'],
          size=9.9,bold=True,align=PP_ALIGN.LEFT)
        t('PFS_MASTER_VALIDATION_PURPOSE_'+str(i),a+w*.40,y,w*.56,vh-.055,
          z['purpose'],size=9.7,align=PP_ALIGN.LEFT)
        if i<len(vals)-1:
            box('PFS_MASTER_VALIDATION_RULE_'+str(i),a+.08,y+vh-.022,w-.16,.012,'border',False)
    # Meaningful editable connections, not chronological arrows between parallel assays.
    line('PFS_MASTER_LINK_MATERIAL_TO_DESIGN',1.86,1.72,2.07,1.72,arrow=True)
    line('PFS_MASTER_LINK_DESIGN_TO_METHODS',4.31,1.72,4.54,1.72,arrow=True)
    line('PFS_MASTER_LINK_METHODS_TO_VALIDATE',8.72,1.72,8.97,1.72,arrow=True)
    # One combined evidence inference beneath methods, bounded against top modules.
    box('PFS_MASTER_SYNTHESIS_BG',.20,3.25,11.43,.99,'light')
    box('PFS_MASTER_SYNTHESIS_ACCENT',.20,3.25,.085,.99,'accent',False)
    t('PFS_MASTER_SYNTHESIS_HEADER',.40,3.40,1.36,.66,'研究结论 / 机制模型',size=12.3,bold=True,color='accent',align=PP_ALIGN.CENTER)
    t('PFS_MASTER_SYNTHESIS_TEXT',1.89,3.36,9.47,.74,m['synthesis'],size=12.4,bold=True,align=PP_ALIGN.LEFT)
    # A convergent connector signals evidence synthesis, not a molecular causality claim.
    line('PFS_MASTER_EVIDENCE_TO_SYNTHESIS',6.61,3.00,6.61,3.21,arrow=True)
    t('PFS_MASTER_NOTE',.28,4.27,11.3,.19,
      m.get('note','注：核心组学并行提供证据，验证实验针对候选机制；图中箭头表示研究设计与证据整合关系。'),
      size=8.2,color='body',align=PP_ALIGN.CENTER)
    # Old framework footer would duplicate the content; blank it.
    foot=next((sh for sh in slide.shapes if sh.name=='PFS_TRANSITION'),None)
    if foot is not None:styled_text(foot,'')


FIGURE_STORY_ROLES = {'context','question','design','analysis','finding','validation','inference','transition'}
FIGURE_STEP_RELATIONS = {'motivates','tests','observes','supports','validates','extends','integrates','parallel','compares','synthesizes'}


def validate_figure_steps(record, slide_i):
    steps=record.get('story_steps')
    if not isinstance(steps,list) or not 2<=len(steps)<=7:
        raise ValueError(f'第{slide_i}页Figure讲述链须提供2–7条story_steps，超出请拆页。')
    for i,step in enumerate(steps):
        if not isinstance(step,dict) or not str(step.get('text','')).strip():
            raise ValueError(f'第{slide_i}页story_steps[{i}]必须有从论文提炼的具体文字。')
        if step.get('role') not in FIGURE_STORY_ROLES:
            raise ValueError(f'第{slide_i}页story_steps[{i}].role无效，不能用笼统方法名占位。')
        if i and step.get('relation_from_previous') not in FIGURE_STEP_RELATIONS:
            raise ValueError(f'第{slide_i}页第{i+1}步必须写relation_from_previous，判断是否真是递进。')
        if 'highlight' in step and str(step['highlight']) not in str(step['text']):
            raise ValueError(f'第{slide_i}页强调词须是本步text内的原词。')
        if step.get('role') in ('finding','validation') and not step.get('panel_ref'):
            raise ValueError(f'第{slide_i}页{step["role"]}步骤必须标注支持它的Fig/panel。')
        if len(str(step['text']))>100:
            raise ValueError(f'第{slide_i}页第{i+1}步太长，需按实验-发现拆成可读短句。')


def figure_story_rich_text(slide,name,x,y,w,h,step,style,fontpt):
    """Plain editable PowerPoint text, optional localized blue emphasis."""
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    sh.name=name
    tf=sh.text_frame;tf.word_wrap=True
    tf.clear();tf.margin_left=tf.margin_right=Inches(.035)
    tf.margin_top=tf.margin_bottom=Inches(.012)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    pa=tf.paragraphs[0];pa.alignment=PP_ALIGN.CENTER
    tx=str(step['text']);em=str(step.get('highlight',''))
    chunks=[(tx,False)]
    if em and em in tx:
        before,after=tx.split(em,1)
        chunks=[(before,False),(em,True),(after,False)]
    for chunk,hi in chunks:
        if not chunk:continue
        run=pa.add_run();run.text=chunk
        run.font.name=CJK_FONT;run.font.size=Pt(fontpt)
        run.font.bold=hi or step.get('role') in ('finding','validation')
        color='accent' if hi or step.get('role') in ('finding','validation') else 'dark'
        run.font.color.rgb=rgb(style.get(color,COLORS[color]))
    return sh


def render_figure_story_steps(slide,record,style=None):
    """Put researcher's authored inference text + arrows BESIDE the original Figure."""
    style=style or COLORS
    card=next((x for x in slide.shapes if x.name in ('PFS_STORY_AREA','DECOR_rightbg')),None)
    if card is None: raise ValueError('模板Figure页缺少PFS_STORY_AREA或DECOR_rightbg放置故事线')
    # Legacy fixed sections are not the author's experimental-story prose. Remove them.
    for shp in list(slide.shapes):
        if shp.name in ('PFS_QUESTION','PFS_DESIGN','PFS_RESULT','PFS_TRANSITION') or shp.name.startswith('DECOR_q_label') or shp.name.startswith('DECOR_d_label') or shp.name.startswith('DECOR_r_label') or shp.name.startswith('DECOR_div_'):
            slide.shapes._spTree.remove(shp.element)
    cx,cy,cw,ch=(z/914400 for z in (card.left,card.top,card.width,card.height))
    x=cx+.20;w=cw-.40;top=cy+.20;available=ch-.40
    steps=record['story_steps'];n=len(steps)
    gap= .24 if n<=5 else .19
    step_h=(available-(n-1)*gap)/n
    if step_h<.35:raise ValueError('Figure故事链过长，建议拆成多个Figure子故事页。')
    for i,st in enumerate(steps):
        ty=top+i*(step_h+gap)
        size=14.6 if n<=3 else (13.5 if n<=5 else 11.8)
        # Fewer words should be large; never shrink everything to illegible fit.
        if len(st['text'])>65: size-=1.4
        figure_story_rich_text(slide,'PFS_FIGURE_STORY_STEP_'+str(i+1),x,ty,w,step_h,st,style,size)
        if i<n-1:
            relation=steps[i+1]['relation_from_previous']
            ay=ty+step_h+.025
            if relation in ('parallel','compares'):
                line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+w*.2), Inches(ay+.08), Inches(x+w*.8), Inches(ay+.08))
                line.name='PFS_FIGURE_STORY_PARALLEL_'+str(i+1)
                line.line.color.rgb=rgb(style['accent']);line.line.width=Pt(1.7)
                line.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            else:
                arrow=slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x+w*.47),Inches(ay),Inches(w*.06),Inches(gap-.035))
                arrow.name='PFS_FIGURE_STORY_ARROW_'+str(i+1)
                arrow.fill.solid();arrow.fill.fore_color.rgb=rgb(style['accent']);arrow.line.fill.background()
    # All reasoning is independently authored, not derived from image similarity.


def validate_argument_graph(record,slide_i):
    nodes=record.get('nodes')
    if not isinstance(nodes,list) or len(nodes)<2:
        raise ValueError(f'第{slide_i}页逻辑图必须至少2个经原文核实的nodes；不允许stages固定链代替。')
    ids=[]
    for n in nodes:
        if not isinstance(n,dict) or not n.get('id') or not n.get('title') or not n.get('detail'):
            raise ValueError(f'第{slide_i}页节点必须包含id/title/detail，title不是技术名词堆叠。')
        if n.get('role','analysis') not in NODE_ROLES:raise ValueError(f'不支持节点role: {n.get("role")}')
        ids.append(n['id'])
    if len(set(ids))!=len(ids):raise ValueError('节点id不得重复')
    edges=record.get('edges',[])
    if not isinstance(edges,list):raise ValueError('edges必须为数组')
    for e in edges:
        if e.get('from') not in ids or e.get('to') not in ids:
            raise ValueError(f'第{slide_i}页edge端点不在nodes中: {e}')
        if e.get('relation') not in EDGE_TYPES:
            raise ValueError(f'edge relation应是{sorted(EDGE_TYPES)}之一')
        if e.get('provenance') not in ('paper','presentation'):
            raise ValueError('每一条科研逻辑关系必须标注provenance=paper或presentation，避免伪造作者因果链')
    if not edges and record.get('diagram_type') not in ('parallel_tracks','comparison','independent_axes'):
        raise ValueError('无连线图必须明确diagram_type是parallel_tracks/comparison/independent_axes')


def populate_slide(slide, record, json_path, index, short_citation, allow_missing):
    kind = record['kind']
    values = {
        'PFS_SECTION': record.get('section', 'RESEARCH PRESENTATION' if kind=='cover' else f'FIGURE STORY  /  {kind.upper()}'),
        'PFS_TITLE': record.get('title',''),
        'PFS_SUBTITLE': record.get('subtitle',''),
        'PFS_CITATION': record.get('citation', short_citation),
        'PFS_NUMBER': f'{index:02d}',
        'PFS_LEAD': record.get('lead', ''),
        'PFS_BULLETS': record.get('bullets', []),
        'PFS_POINTS': record.get('points', []),
        'PFS_QUESTION': record.get('question', ''),
        'PFS_DESIGN': record.get('design', ''),
        'PFS_RESULT': ('\n'.join(filter(None,[record.get('result',''), record.get('takeaway','')]))),
        'PFS_TRANSITION': record.get('transition', ''),
        'PFS_FIGURE_REF': record.get('figure_ref',''),
        'PFS_LIMITATION': record.get('limitation',''),
    }
    for shape in list(slide.shapes):
        if shape.name in values:
            val = values[shape.name]
            styled_text(shape, val, is_list=shape.name in ('PFS_BULLETS', 'PFS_POINTS'))
    if kind == 'figure' and record.get('story_steps'):
        render_figure_story_steps(slide,record,style=getattr(slide,'_pfs_style',COLORS))
    if kind == 'figure':
        img_shape = next((v for v in slide.shapes if v.name == 'PFS_IMAGE'), None)
        if img_shape is None:
            raise ValueError('模板figure页缺少 PFS_IMAGE 插槽')
        image_paths = record.get('images', []) or ([record['image']] if record.get('image') else [])
        insert_images(slide, img_shape, image_paths, json_path, allow_missing)
    if kind == 'framework':
        if record.get('study_map'):
            render_study_map(slide,record,style=getattr(slide,'_pfs_style',COLORS))
        else:
            render_integrated_framework(slide,record,style=getattr(slide,'_pfs_style',COLORS))
    if kind in ('logic','route'):
        render_argument_graph(slide,record,style=getattr(slide,'_pfs_style',COLORS))
    # Human speaking layer; the reviewed page JSON remains the machine record.
    note=['本页问题：'+record.get('question',record.get('title',''))]
    if record.get('lead'):note.append('背景与边界：'+str(record['lead']))
    if record.get('story_steps'):
        note.append('讲述顺序：\n'+'\n'.join(str(i+1)+'．'+str(s['text']) for i,s in enumerate(record['story_steps'])))
    elif record.get('bullets'):note.append('讲述要点：\n'+'\n'.join(map(str,record['bullets'])))
    if record.get('limitation'):note.append('证据边界：'+str(record['limitation']))
    if record.get('transition'):note.append('衔接：'+str(record['transition']))
    if isinstance(record.get('speaker_notes'),str):note.append(record['speaker_notes'])
    note.append('来源：'+record.get('figure_ref',record.get('citation',short_citation)))
    slide.notes_slide.notes_text_frame.text='\n\n'.join(note)


def check_input(manifest):
    if not isinstance(manifest.get('slides'), list) or not manifest['slides']:
        raise ValueError('JSON必须有非空slides数组')
    for i, item in enumerate(manifest['slides'], start=1):
        if item.get('kind') not in VALID_KINDS: raise ValueError(f'第{i}页kind无效: {item.get("kind")}')
        if not str(item.get('title','')).strip(): raise ValueError(f'第{i}页缺少title')
        if item['kind']=='framework':
            if item.get('study_map'):validate_study_map(item,i)
            else:validate_framework(item,i)
        if item['kind'] in ('logic','route'):
            validate_argument_graph(item,i)
        if item['kind']=='figure':
            if not str(item.get('figure_ref','')).strip():
                raise ValueError(f'第{i}页Figure故事缺少 figure_ref')
            if item.get('story_steps'):
                validate_figure_steps(item,i)
            else:
                for need in ('question','design','result','transition'):
                    if not str(item.get(need,'')).strip():
                        raise ValueError(f'第{i}页旧版Figure说明缺少 {need}；推荐改用story_steps讲述链。')


def build(json_path, out_path, template, mode, allow_missing):
    if Image is None:
        raise RuntimeError('缺少 Pillow：pip install Pillow')
    manifest = json.loads(json_path.read_text(encoding='utf-8'))
    check_input(manifest)
    template_path = Path(template) if template else DEFAULT_PATH
    if not template_path.exists() and not template:
        make_default_template(DEFAULT_PATH)
    if not template_path.is_file(): raise FileNotFoundError(f'模板不存在：{template_path}')
    if template_path.suffix.lower() != '.pptx':
        raise ValueError('生成器目前接收.pptx模板；.potx/.ppt请先转换为.pptx并检查格式')
    prs = Presentation(str(template_path))
    archetypes = find_archetypes(prs)
    resolved_mode = 'tagged' if (mode == 'auto' and archetypes) else ('native' if mode=='auto' else mode)
    shortcite = manifest.get('meta',{}).get('short_citation','来源：原论文')
    original = list(prs.slides)
    original_ids = list(prs.slides._sldIdLst)
    W = prs.slide_width / 914400; H = prs.slide_height / 914400
    warnings=[]
    if resolved_mode=='tagged':
        if not archetypes: raise ValueError('模板无PFS_LAYOUT_*标记，改用--template-mode native或制作标记模板')
        tagged_style=theme_rgb_from_shape(prs) if template else COLORS
        for idx, rec in enumerate(manifest['slides'], start=1):
            exemplar=archetypes.get(rec['kind']) or archetypes.get('text')
            if exemplar is None:
                raise ValueError(f'模板缺少{rec["kind"]}版式且没有text备用版式')
            slide=duplicate_slide(prs, exemplar)
            slide._pfs_style=tagged_style
            populate_slide(slide,rec,json_path,idx,shortcite,allow_missing)
    elif resolved_mode=='native':
        warnings.append('普通模板采用母版/版式及配色采样自动适配；若需完全复刻，请标注内容槽位并使用tagged模式。')
        style = theme_rgb_from_shape(prs)
        blank = next((lay for lay in prs.slide_layouts if 'blank' in lay.name.lower() or '空白' in lay.name), prs.slide_layouts[-1])
        for idx,rec in enumerate(manifest['slides'],start=1):
            slide=prs.slides.add_slide(blank)
            for s in list(slide.shapes): slide.shapes._spTree.remove(s.element)
            add_frame(slide,rec['kind'],W,H,style=style)
            slide._pfs_style=style
            populate_slide(slide,rec,json_path,idx,shortcite,allow_missing)
    else:
        raise ValueError('模板模式仅支持auto/tagged/native')
    for old_id in original_ids: remove_slide(prs,old_id)
    out_path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(str(out_path))
    from apply_script_fonts import apply_fonts
    apply_fonts(out_path)
    print(f'PPTX: {out_path} | slides={len(prs.slides)} | mode={resolved_mode} | template={template_path.name}')
    for w in warnings: print('WARN:',w)
    return out_path


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('data',nargs='?',type=Path,help='已核实的06_PPT页面数据.json')
    parser.add_argument('--make-default-template',type=Path)
    parser.add_argument('--template',type=Path,help='优先使用用户上传的.pptx模板')
    parser.add_argument('--template-mode',choices=['auto','tagged','native'],default='auto')
    parser.add_argument('--output',type=Path)
    parser.add_argument('--allow-missing-images',action='store_true',help='仅限演示/草稿')
    a=parser.parse_args()
    if a.make_default_template:
        print('Created:', make_default_template(a.make_default_template)); return 0
    if not a.data or not a.output:
        parser.error('生成PPTX需要data JSON与--output；或用--make-default-template')
    try: build(a.data.resolve(),a.output.resolve(),a.template,a.template_mode,a.allow_missing_images)
    except Exception as e:
        print('ERROR:',e,file=sys.stderr);return 1
    return 0


if __name__=='__main__':
    sys.exit(main())
