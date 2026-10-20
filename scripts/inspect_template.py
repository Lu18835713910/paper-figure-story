#!/usr/bin/env python3
"""Inspect a PowerPoint template: dimensions, layouts, tagged slots, slide-level shape positions."""
import argparse
import json
from pathlib import Path
from pptx import Presentation


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('template',type=Path)
    p.add_argument('--output',type=Path,help='save complete JSON inspection (default stdout)')
    p.add_argument('--max-slides',type=int,default=30)
    a=p.parse_args()
    prs=Presentation(str(a.template))
    obj={
      'template':a.template.name,
      'ratio':round(prs.slide_width/prs.slide_height,5),
      'width_in':round(prs.slide_width/914400,3),
      'height_in':round(prs.slide_height/914400,3),
      'slide_count':len(prs.slides),
      'masters':len(prs.slide_masters),
      'layouts':[{'index':i,'name':x.name,'placeholders':[{'name':s.name,'type':str(s.placeholder_format.type)} for s in x.placeholders]} for i,x in enumerate(prs.slide_layouts)],
      'slides':[],
    }
    for i,slide in enumerate(list(prs.slides)[:a.max_slides],1):
        shapes=[]
        for s in slide.shapes:
            if s.name.startswith('PFS_') or s.has_text_frame or s.shape_type in (13,):
                desc={'name':s.name,'type':str(s.shape_type),
                  'x':round(s.left/914400,2),'y':round(s.top/914400,2),
                  'w':round(s.width/914400,2),'h':round(s.height/914400,2)}
                if s.has_text_frame:
                    desc['text_preview']=s.text.replace('\n',' ')[:100]
                shapes.append(desc)
        tags=[s.name for s in slide.shapes if s.name.startswith('PFS_LAYOUT_')]
        obj['slides'].append({'index':i,'layout':slide.slide_layout.name,'tagged_kind':tags,'shapes':shapes})
    raw=json.dumps(obj,ensure_ascii=False,indent=2)
    if a.output:
        a.output.parent.mkdir(parents=True,exist_ok=True)
        a.output.write_text(raw,encoding='utf-8')
        print('Saved',a.output)
    else: print(raw)

if __name__=='__main__':main()
