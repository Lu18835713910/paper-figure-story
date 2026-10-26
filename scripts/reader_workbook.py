"""Reader-oriented formatting of existing, source-bound research logic and Methods.

This is a document layout, not a scientific inference or extraction algorithm.
"""
from pathlib import Path
import csv,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START
from PIL import Image
import research_logic as rl
from build_research_workbook import fonts,table,hyperlink,page_setup,value
from apply_script_fonts import apply_docx_fonts

def build_reader(model_path,manifest_path,output,annotations_path,sample=False):
 m=rl.load(model_path);validation=rl.validate_model(m)
 if not validation['valid']:raise ValueError('READER_MODEL_INVALID')
 a=rl.load(annotations_path)
 if a.get('logic_snapshot')!=m['snapshot'] or a.get('formal_ready') is not False:raise ValueError('READER_VERSION_OR_REVIEW')
 methods=rl.methods_view(m);em={e['experiment_id']:e for e in methods};nm={n['id']:n for n in m['nodes']};edges={e['id']:e for e in m['edges']}
 for r in a['rows']:
  if not r.get('source_ids') or not set(r['source_ids'])<=set(m['sources']):raise ValueError('READER_SOURCE')
 for eid,card in a['cards'].items():
  if eid not in em or not card.get('source_ids') or not set(card['source_ids'])<=set(m['sources']):raise ValueError('READER_CARD_SOURCE')
  if not set(card['field_bindings'])<=set(em[eid]['fields']):raise ValueError('READER_CARD_FIELD')
 allids=[eid for g in a['groups'] for eid in g['experiment_ids']]
 if len(allids)!=len(set(allids)) or set(allids)!=set(em):raise ValueError('READER_METHODS_COVERAGE')
 graphs={g['view']:g for g in rl.load(manifest_path)['graphs']}
 wanted=['A','interface'] if sample else ['A','B']+[g['view'] for g in a['groups'] if g.get('view')]
 for view in wanted:
  g=graphs[view]
  if g['logic_snapshot']!=m['snapshot'] or rl.file_sha(g['pptx'])!=g['pptx_sha256'] or rl.file_sha(g['png'])!=g['png_sha256']:raise ValueError('READER_GRAPH_DRIFT')
 out=Path(output)
 if out.exists():raise ValueError('OUTPUT_EXISTS')
 out.parent.mkdir(parents=True,exist_ok=True);doc=Document();page_setup(doc.sections[0])
 doc.core_properties.title=a['title'];doc.core_properties.identifier=m['snapshot'];doc.core_properties.author='Paper Figure Story'
 for name,size in [('Normal',11),('Title',27),('Heading 1',19),('Heading 2',14),('Caption',9.5)]:
  st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=name.startswith('Heading');st.font.color.rgb=RGBColor.from_string('174994' if name.startswith('Heading') else '111827');st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.06
  st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 f=doc.sections[0].footer.paragraphs[0];fonts(f.add_run(a['short_title']+'   '),9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 used=[];shown=[];fresh_section=[False]
 def p(t,style=None):return doc.add_paragraph(t,style)
 def new(t):
  if not fresh_section[0]:doc.add_page_break()
  fresh_section[0]=False;doc.add_heading(t,1)
 def source(sids):return '原文 p'+', '.join(str(v) for v in sorted({m['sources'][sid]['physical_page'] for sid in sids}))
 def diagram(view,maxw=10.4,maxh=4.9):
  g=graphs[view];im=Image.open(g['png']);width=min(maxw,maxh*im.width/im.height);r=p('').add_run();r.add_picture(g['png'],width=Inches(width));used.append(g)
  p('此图与配套可编辑PPT使用相同科学节点和关系。','Caption')
 def teaching(key):
  row=next(r for r in a['rows'] if r.get('key')==key);doc.add_heading(row['title'],2);p(row['text']);p(source(row['source_ids']),'Caption')
 def branch(group):
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec,group.get('page_width',16.54),group.get('page_height',11.69));doc.add_heading(group['title'],1)
  diagram(group['view'],group.get('page_width',16.54)-1.2,group.get('diagram_height',7.5));teaching(group['teaching_key'])
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec)
  fresh_section[0]=True
 def cards(eids,title):
  new(title);t=doc.add_table(rows=1,cols=2);t.autofit=False
  cw=5.2
  for cell,eid in zip(t.rows[0].cells,eids):
   cell.width=Inches(cw);r=em[eid];c=a['cards'][eid];shown.append(eid)
   head=cell.paragraphs[0];head.paragraph_format.space_after=Pt(7);fonts(head.add_run(c.get('title',r['name'])),14,True,'174994')
   plain=r['plain_explanation']
   entries=[('实验目的',plain['why']),('研究对象',c['object']),('比较与对照',c['groups']),('关键处理',c['steps']),('检测指标',plain['readout']),('主要结果',plain['result']),('科学意义',plain['meaning']),('不能据此证明',c.get('limit',plain['limit']))]
   for label,txt in entries:
    para=cell.add_paragraph();para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1.05;fonts(para.add_run(label+'  '),10.5,True);fonts(para.add_run(txt),10.5)
   if c.get('teaching_key') and len(eids)>1:
    tip=next(r for r in a['rows'] if r.get('key')==c['teaching_key']);para=cell.add_paragraph();fonts(para.add_run('读图提示  '),10.5,True,'174994');fonts(para.add_run(tip['text']),10.5)
   para=cell.add_paragraph();fonts(para.add_run(source(c['source_ids'])+'  '+', '.join(f.replace('CNNM4F','Fig.') for f in r['figure_ids'])),9,color='61748A')
  if len(eids)==1:
   cell=t.rows[0].cells[1];cell.width=Inches(cw);c=a['cards'][eids[0]]
   tip=next((r for r in a['rows'] if r.get('key')==c.get('teaching_key')),None)
   title=tip['title'] if tip else c.get('companion_title','怎样把这个结果放回研究主线')
   fonts(cell.paragraphs[0].add_run(title),14,True,'174994')
   for txt in ([tip['text']] if tip else c.get('companion_texts',[em[eids[0]]['plain_explanation']['meaning']])):
    para=cell.add_paragraph();para.paragraph_format.space_after=Pt(8);fonts(para.add_run(txt),11)
   para=cell.add_paragraph();fonts(para.add_run(source(c['source_ids'])),9,color='61748A')
  t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 doc.add_heading(a['title'],0);p(a['subtitle']);doc.add_heading('第一部分 三分钟看懂这篇论文',1)
 opening=[r for r in a['rows'] if r['kind']=='opening'];table(doc,['你可能先问','这篇论文给出的回答'],[[r['title'],r['text']] for r in opening],[2.0,8.4],11.5)
 p('阅读顺序：先看研究主线，再按分支理解设计和结果；完整参数、编号和审核状态统一放在附录。')
 new('用一张简图把握研究主线');diagram('A',10.4,5.55)
 if sample:
  branch(a['sample_branch'])
  cards([a['sample_experiment']],a['sample_title'])
 else:
  new('第二部分 作者怎样逐层追问')
  stages=[r for r in a['rows'] if r['kind']=='stage']
  table(doc,['已经知道什么','为什么还要继续','下一步获得什么认识'],[[r['known'],r['title'],r['text']] for r in stages],[3.1,3.6,3.7],11)
  p('以下进入各分支的具体实验设计。每个实验卡保留对象、比较、处理、读数、结果和科学边界。')
  for group in a['groups']:
   if group.get('view'):
    branch(group)
   for k in range(0,len(group['experiment_ids']),2):cards(group['experiment_ids'][k:k+2],'第三部分 '+group['card_title'])
  # A poster page keeps detailed-node body type readable at its native scale.
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec,20,17);doc.add_heading('第四部分 不同实验怎样互相支持',1);diagram('B',18.5,14.7)
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec);doc.add_heading('把证据与解释分开读',1)
  relation_rows=[]
  for rec in a['relationships']:
   e=edges[rec['edge_id']];relation_rows.append([rec['link'],rec['meaning'],rec['limit'],source(e['source_ids'])])
  table(doc,['哪些证据连接','这条关系说明什么','不能额外推出','来源'],relation_rows,[2.25,3.4,3.8,.95],10.5)
  new('第五部分 完整Methods与来源附录')
  p('这里保留查询用的任务编号。'+str(len(methods))+'项包含实验、控制、计算及共用流程，不等于同等数量的独立实验。完整'+str(sum(len(r['fields']) for r in methods))+'字段未压缩删除；它们在配套原记录和CSV中逐项保留。')
  for row in a['rows']:
   if row['kind']=='method_note':p(row['text'])
  invpath=out.parent/'Full_Experimental_Inventory.source.json';shutil.copy2(m['inputs']['inventory']['path'],invpath)
  csvpath=out.parent/'Methods字段与来源完整索引.csv'
  with csvpath.open('w',encoding='utf-8-sig',newline='') as f:
   writer=csv.writer(f);writer.writerow(['实验任务ID','字段','原值','原核验状态','原文来源','限制'])
   for task in methods:
    for key,field in task['fields'].items():writer.writerow([task['experiment_id'],key,value(field.get('value')),field.get('status'),json.dumps(field.get('source_refs',[]),ensure_ascii=False),field.get('reason','')])
  for j in range(0,len(methods),9):
   if j:new('Methods任务索引续表')
   table(doc,['任务编号','研究或分析任务','类型','对应图','核验说明'],[[r['experiment_id'],r['name'],{'experiment':'实验','control_experiment':'控制','computational_analysis':'计算分析','shared_technical_preparation':'共用制备','shared_computational_workflow':'共用计算'}[r['task_type']],','.join(f.replace('CNNM4F','Fig.') for f in r['figure_ids']),'完整已知字段及未核项见配套记录'] for r in methods[j:j+9]],[1.45,4.5,1.0,1.1,2.35],10)
  q=p('完整Methods原记录  ');hyperlink(q,invpath.name,invpath.name);q=p('逐字段来源与状态  ');hyperlink(q,csvpath.name,csvpath.name)
  new('哪些信息仍需回查')
  statuses=[r for r in a['rows'] if r['kind']=='status'];table(doc,['待核事项','对理解的影响'],[[r['title'],r['text']] for r in statuses],[2.0,8.4],11)
  terms=[r for r in a['rows'] if r['kind']=='glossary'];doc.add_heading('术语速查',2);table(doc,['名称','它是什么'],[[r['title'],r['text']] for r in terms],[2.0,8.4],11)
  new('技术附录 版本和原文定位')
  p('以下为复核资料，不需要背记。文件及摘录定位成功不等于独立科学终审通过。来源语义审阅尚未全部完成，用户对本版可读性的认可也未取得。')
  p('逻辑版本 '+m['snapshot'],'Caption');p('原Inventory SHA256 '+rl.file_sha(m['inputs']['inventory']['path']),'Caption')
  table(doc,['来源索引','论文物理页','资料层次'],[[sid,str(src['physical_page']),'正文、图注或Methods；定位与状态保留在配套审计中'] for sid,src in m['sources'].items()],[2.0,1.4,7.0],9.5)
 for root in [doc.element,doc.styles.element]:
  for b in list(root.iter(qn('w:pBdr'))):b.getparent().remove(b)
 doc.save(out);apply_docx_fonts(out)
 rec=dict(version='pfs-reader-workbook-1',docx=str(out),docx_sha256=rl.file_sha(out),logic_snapshot=m['snapshot'],network_snapshot=m['network_snapshot'],annotations_sha256=rl.file_sha(annotations_path),graphs=used,experiment_ids=shown,inventory_task_ids=list(em),inventory_fields=sum(len(r['fields']) for r in methods),sample=sample,document_text_tables_editable=True,diagram_editing='Linked native PPTX; identical exported images embedded in DOCX',formal_ready=False,USER_ACCEPTANCE='PENDING')
 out.with_suffix('.receipt.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2),'utf-8');return rec
