"""Synthetic, local-only Visual Polish safety and interface regression tests.

No paper fixtures, Office renderer, PDF extraction or network calls are needed.
Actual Office and visual checks are separately recorded for the real-paper runs.
"""
from pathlib import Path
import copy,json,subprocess,sys,tempfile,unittest
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile,ZIP_DEFLATED
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import polish_presentation as v
import run_pfs
from pptx import Presentation
from pptx.util import Inches,Pt
from PIL import Image

class VisualPolish(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.src=self.root/'source.pptx'
  p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
  for n in range(2):
   s=p.slides.add_slide(p.slide_layouts[6]);q=s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2));q.name='CONTENT';q.text='WT compared with mutant; residual 80% at 5 minutes suggests weaker efflux.' if n==0 else 'Important scientific limitation remains PARTIAL.';q.text_frame.paragraphs[0].font.size=Pt(18)
   s.notes_slide.notes_text_frame.text='Source details and controls; not a verified causal mechanism.'
  im=self.root/'figure.png';Image.new('RGB',(120,80),'white').save(im);q=p.slides[0].shapes.add_picture(str(im),Inches(1),Inches(4),width=Inches(3));self.pic=q.shape_id
  self.image=im;p.save(self.src);self.sha=v.digest(self.src);self.data=v.package(self.src);self.parts=v.active_parts(self.data);self.sid=2
 def tearDown(self):self.tmp.cleanup()
 def plan(self,actions=None):return dict(version=v.VERSION,input_sha256=self.sha,actions=actions or [])
 def run_polish(self,plan=None,context=None,**kw):return v.polish(self.src,self.root/'out',renderer='none',plan=plan,context=context,**kw)
 def text(self,t):return dict(text=t,refs=[dict(kind='deck',slide=1)])
 def test_01_ast_imports(self):self.assertEqual(len(self.parts),2)
 def test_02_preserves_original_and_media(self):
  r=self.run_polish();self.assertEqual(v.digest(self.src),self.sha);self.assertEqual(v.digest(self.root/'out/original.pptx'),self.sha);self.assertTrue(r['invariants']['original_media_parts_unchanged'])
 def test_03_no_renderer_no_visual_pass(self):
  r=self.run_polish();self.assertEqual(r['status'],'PARTIAL');self.assertFalse(r['actual_powerpoint_export']);self.assertEqual(r['USER_ACCEPTANCE'],'PENDING')
 def test_04_actual_font_properties(self):
  self.run_polish();self.assertTrue(v.check_fonts(self.root/'out/polished.pptx')['valid'])
 def test_05_output_overwrite_refused(self):
  self.run_polish()
  with self.assertRaisesRegex(v.PolishError,'OUTPUT_EXISTS'):self.run_polish()
 def test_06_third_cycle_rejected(self):
  with self.assertRaisesRegex(v.PolishError,'PROFILE_OR_CYCLE'):self.run_polish(cycle=3)
 def test_07_plan_hash_drift(self):
  p=self.plan();p['input_sha256']='0'*64
  with self.assertRaisesRegex(v.PolishError,'PLAN_VERSION_OR_INPUT_DRIFT'):self.run_polish(p)
 def test_08_unsourced_text_rejected(self):
  with self.assertRaisesRegex(v.PolishError,'UNBOUND_TEXT'):v.bound_text({'text':'invented'},self.data,self.parts,{})
 def test_09_reversed_result_rejected(self):
  with self.assertRaisesRegex(v.PolishError,'TEXT_NOT_IN_BOUND_SOURCE'):v.bound_text(self.text('stronger efflux'),self.data,self.parts,{})
 def test_10_changed_number_rejected(self):
  with self.assertRaisesRegex(v.PolishError,'TEXT_NOT_IN_BOUND_SOURCE'):v.bound_text(self.text('90% at 5 minutes'),self.data,self.parts,{})
 def test_11_unadjudicated_upstream_refused(self):
  with self.assertRaisesRegex(v.PolishError,'UPSTREAM_NOT_ADJUDICATED'):v.bound_text({'text':'new','refs':[{'kind':'upstream','response':'r'}]},self.data,self.parts,{})
 def test_12_critical_region_crop_detected(self):
  p=Presentation(self.src);p.slides[0].shapes[-1].crop_bottom=.3;p.save(self.src)
  a=dict(id='asset',path=str(self.image),sha256=v.digest(self.image),panels=['F1a'],panel_mapping_status='CHECKED',critical_regions=[dict(id='axis',bbox_px=[0,0,120,80])])
  d=v.audit_input(self.src,{'assets':[a]});self.assertIn('CRITICAL_REGION_CLIPPED',[x['code'] for x in d['slides'][0]['issues']])
 def test_13_all_not_fine_panel_pass(self):
  a=dict(sha256=v.digest(self.image),panels=['all'],panel_mapping_status='CHECKED');r=v.audit_input(self.src,dict(assets=[a],slides=[dict(slide=1,expected_panels=['F1a'])]));self.assertIn('PANEL_MAPPING_UNRESOLVED',[x['code'] for x in r['slides'][0]['issues']])
 def test_14_mapping_unreviewed_not_pass(self):
  a=dict(sha256=v.digest(self.image),panels=['F1a'],panel_mapping_status='PARTIAL');r=v.audit_input(self.src,dict(assets=[a],slides=[dict(slide=1,expected_panels=['F1a'])]));self.assertIn('PANEL_MAPPING_UNRESOLVED',[x['code'] for x in r['slides'][0]['issues']])
 def test_15_coverage_cannot_be_closed_by_layout(self):
  ctx={'slides':[dict(slide=1,required_content=[dict(id='gap',phrases=['missing comparison'])])]};p=self.plan([dict(id='move',kind='layout',slide=1,shapes=[dict(shape_id=2,position=[1,1,8,2])])]);p['resolves']={'1:UPSTREAM_CONTENT_REQUIRED':['move']};r=self.run_polish(p,ctx);self.assertTrue(any(i['code']=='UPSTREAM_CONTENT_REQUIRED' for i in r['unresolved']))
 def test_16_arrow_identity_not_upgraded(self):
  a=dict(id='edge',kind='narrative',slide=1,identity='causal',relation_id='r',steps=[],arrows=[]);r=self.run_polish(self.plan([a]),dict(relationships=[dict(id='r',identity='presentation')]));self.assertIn('ARROW_IDENTITY_UPGRADE',r['failed_actions'][0]['reason'])
 def test_17_unregistered_scientific_arrow_refused(self):
  a=dict(id='edge',kind='narrative',slide=1,identity='scientific',steps=[],arrows=[]);r=self.run_polish(self.plan([a]));self.assertIn('UNREGISTERED',r['failed_actions'][0]['reason'])
 def test_18_native_table_and_notes(self):
  a=dict(id='tab',kind='table',slide=1,remove_shapes=[2],headers=['条件','观察结果'],rows=[[self.text('WT'),self.text('residual 80%')]],position=[1,1,10,2],column_widths=[3,7],font_size=16)
  r=self.run_polish(self.plan([a]));p=Presentation(r['output']);self.assertTrue(any(q.has_table for q in p.slides[0].shapes));self.assertFalse(r['invariants']['unintended_text_loss']);self.assertIn('weaker efflux',p.slides[0].notes_slide.notes_text_frame.text)
 def test_19_backup_cannot_hide_important_limit(self):
  ctx={'slides':[dict(slide=2,important_limits=[dict(text='Important scientific limitation remains PARTIAL.',refs=[dict(kind='deck',slide=2)])])]};a=dict(id='hide',kind='backup',slide=2);r=self.run_polish(self.plan([a]),ctx);self.assertIn('SCIENTIFIC_LIMIT_WOULD_BE_HIDDEN',r['failed_actions'][0]['reason'])
 def test_20_protected_source_drift(self):
  with self.assertRaisesRegex(v.PolishError,'PROTECTED_SOURCE_DRIFT'):self.run_polish(context={'protected_sources':[dict(path=str(self.src),sha256='x')]})
 def test_21_picture_restore_requires_review(self):
  a=dict(id='pic',kind='replace_image',slide=1,shape_id=self.pic,asset_id='a',required_panels=['F1a'],position=[1,4,3,2]);ctx={'assets':[dict(id='a',path=str(self.image),sha256=v.digest(self.image),panels=['F1a'],panel_mapping_status='PARTIAL')]};r=self.run_polish(self.plan([a]),ctx);self.assertIn('IMAGE_NOT_SOURCE_REVIEWED',r['failed_actions'][0]['reason'])
 def test_22_replacement_preserves_original_wording_in_notes(self):
  old='WT compared with mutant; residual 80% at 5 minutes suggests weaker efflux.';a=dict(id='txt',kind='replace_text',slide=1,shape_id=2,expected_text=old,replacement=self.text('residual 80% at 5 minutes'),position=[1,1,8,1]);r=self.run_polish(self.plan([a]));self.assertFalse(r['invariants']['unintended_text_loss'])
 def test_23_auto_off_never_calls_polish(self):
  with patch.object(v,'polish') as mock:
   r=run_pfs.postprocess_main(SimpleNamespace(no_polish=True),self.root,'balanced');self.assertEqual(r['status'],'DISABLED_BY_OPTION');mock.assert_not_called()
 def test_24_auto_hook_uses_finished_ppt_only(self):
  p=self.root/'pptx';p.mkdir();(p/'main_balanced.pptx').write_bytes(self.src.read_bytes())
  with patch.object(v,'polish',return_value=dict(status='PARTIAL',output='finished')) as mock:
   r=run_pfs.postprocess_main(SimpleNamespace(),self.root,'balanced');self.assertEqual(r['scientific_review'],'INHERITED_NOT_UPGRADED');self.assertEqual(mock.call_args.args[0],p/'main_balanced.pptx')
 def test_25_empty_preview_receipt_not_valid(self):
  receipt=self.root/'receipt.json';v.save(receipt,dict(input_sha256=self.sha,source_unchanged=True,rendered=[]))
  with self.assertRaisesRegex(v.PolishError,'EMPTY_RENDER_RECEIPT'):v.render(self.src,self.root/'preview','balanced',provided=receipt)
 def test_26_real_cli_polish_route(self):
  r=subprocess.run([sys.executable,str(Path(run_pfs.__file__)),'polish','--pptx',str(self.src),'--outdir',str(self.root/'cli'),'--renderer','none'],capture_output=True,text=True,encoding='utf-8');self.assertEqual(r.returncode,0,r.stderr);self.assertEqual(v.read(self.root/'cli/visual_polish_report.json')['science_audit'],'NOT_PERFORMED_BY_VISUAL_POLISH')
 def test_27_modes_share_science_requirements(self):
  self.assertEqual(v.MODES,dict(quick=1280,balanced=1920,full=2560))
 def test_28_absolute_relationships_resolve(self):self.assertEqual(v.resolve_target('ppt/slides/s.xml','/ppt/notesSlides/n.xml'),'ppt/notesSlides/n.xml')
 def test_29_non_input_never_counts_as_user_acceptance(self):
  r=self.run_polish();self.assertEqual(r['scientific_review'],'NOT_AVAILABLE');self.assertEqual(r['USER_ACCEPTANCE'],'PENDING')
 def test_30_insert_uses_original_style_not_previous_table(self):
  a=dict(id='a',kind='table',slide=1,remove_shapes=[2],headers=['条件'],rows=[[self.text('WT')]],position=[1,1,8,2],font_size=16)
  b=dict(id='b',kind='insert_table_slide',slide=1,after=1,clear_names=['CONTENT'],headers=['条件'],rows=[[self.text('mutant')]],position=[1,1,8,2],font_size=16)
  r=self.run_polish(self.plan([a,b]));p=Presentation(r['output']);self.assertEqual(len(p.slides),3);self.assertEqual(sum(q.has_table for q in p.slides[1].shapes),1)
 def test_31_title_claim_without_mapping_is_unresolved(self):
  p=Presentation(self.src);p.slides[0].shapes[0].text='Figure 3A–C';p.save(self.src);r=v.audit_input(self.src);issue=next(x for x in r['slides'][0]['issues'] if x['code']=='PANEL_MAPPING_UNRESOLVED');self.assertEqual(issue['panels'],['3A','3B','3C'])

if __name__=='__main__':unittest.main(verbosity=2)
